from __future__ import annotations

import json
import re
import sqlite3
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from moysklad_price import MoySkladClient, _canonical_href


ACTIVITY_START = datetime(2026, 4, 1)


@dataclass(frozen=True)
class ActivityReport:
    week: str
    generated_at: str
    period_start: str
    period_end: str
    common_path: str
    summary: dict
    personal_paths: dict[str, str]
    sent_at: str | None


def client_parts(agent_name: str) -> tuple[str, str]:
    display = (agent_name.split("/", 1)[0] or agent_name).strip()
    key = re.sub(r"[^\w]+", " ", display.replace("_", " ").casefold().replace("ё", "е"))
    return " ".join(key.split()), display


def _href(value: object) -> str:
    if not isinstance(value, dict):
        return ""
    return _canonical_href(str(value.get("meta", {}).get("href", "")))


def _name(value: object) -> str:
    return str(value.get("name", "")).strip() if isinstance(value, dict) else ""


class ActivityDatabase:
    def __init__(self, path: Path):
        self.path = path
        path.parent.mkdir(parents=True, exist_ok=True)
        with self._connect() as db:
            db.executescript(
                """
                CREATE TABLE IF NOT EXISTS shipments (
                    demand_id TEXT PRIMARY KEY,
                    demand_name TEXT NOT NULL DEFAULT '',
                    moment TEXT NOT NULL,
                    amount_minor INTEGER NOT NULL DEFAULT 0,
                    agent_href TEXT NOT NULL DEFAULT '',
                    agent_name TEXT NOT NULL DEFAULT '',
                    client_key TEXT NOT NULL,
                    client_name TEXT NOT NULL,
                    channel_href TEXT NOT NULL DEFAULT '',
                    channel_name TEXT NOT NULL DEFAULT '',
                    updated TEXT NOT NULL DEFAULT ''
                );
                CREATE INDEX IF NOT EXISTS shipments_moment_idx ON shipments(moment);
                CREATE INDEX IF NOT EXISTS shipments_client_idx ON shipments(client_key, moment);
                CREATE TABLE IF NOT EXISTS activity_meta (
                    key TEXT PRIMARY KEY,
                    value TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS activity_reports (
                    week TEXT PRIMARY KEY,
                    generated_at TEXT NOT NULL,
                    period_start TEXT NOT NULL,
                    period_end TEXT NOT NULL,
                    common_path TEXT NOT NULL,
                    summary_json TEXT NOT NULL,
                    personal_paths_json TEXT NOT NULL,
                    sent_at TEXT
                );
                CREATE TABLE IF NOT EXISTS activity_deliveries (
                    week TEXT NOT NULL,
                    user_id INTEGER NOT NULL,
                    sent_at TEXT NOT NULL,
                    PRIMARY KEY(week, user_id)
                );
                """
            )

    def _connect(self):
        db = sqlite3.connect(self.path, timeout=30)
        db.row_factory = sqlite3.Row
        return db

    def backup_to(self, destination: Path) -> None:
        destination.parent.mkdir(parents=True, exist_ok=True)
        with self._connect() as source, sqlite3.connect(destination) as target:
            source.backup(target)

    def meta(self, key: str) -> str | None:
        with self._connect() as db:
            row = db.execute("SELECT value FROM activity_meta WHERE key=?", (key,)).fetchone()
        return row["value"] if row else None

    def set_meta(self, key: str, value: str) -> None:
        with self._connect() as db:
            db.execute(
                "INSERT INTO activity_meta(key,value) VALUES(?,?) "
                "ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, value)
            )

    def upsert_shipments(self, rows: list[dict]) -> int:
        values = []
        for row in rows:
            demand_id = str(row.get("id") or _href(row)).strip()
            moment = str(row.get("moment") or "").strip()
            agent = row.get("agent")
            agent_name = _name(agent)
            key, display = client_parts(agent_name)
            if not demand_id or not moment or not key:
                continue
            channel = row.get("salesChannel")
            values.append((
                demand_id, str(row.get("name") or ""), moment,
                int(Decimal(str(row.get("sum") or 0))), _href(agent), agent_name,
                key, display, _href(channel), _name(channel), str(row.get("updated") or ""),
            ))
        with self._connect() as db:
            db.executemany(
                """INSERT INTO shipments VALUES(?,?,?,?,?,?,?,?,?,?,?)
                ON CONFLICT(demand_id) DO UPDATE SET
                demand_name=excluded.demand_name,moment=excluded.moment,
                amount_minor=excluded.amount_minor,agent_href=excluded.agent_href,
                agent_name=excluded.agent_name,client_key=excluded.client_key,
                client_name=excluded.client_name,channel_href=excluded.channel_href,
                channel_name=excluded.channel_name,updated=excluded.updated""", values
            )
        return len(values)

    def shipments(self) -> list[sqlite3.Row]:
        with self._connect() as db:
            return db.execute("SELECT * FROM shipments ORDER BY moment").fetchall()

    def shipment_count(self) -> int:
        with self._connect() as db:
            return int(db.execute("SELECT COUNT(*) FROM shipments").fetchone()[0])

    def save_report(self, report: ActivityReport) -> None:
        with self._connect() as db:
            db.execute(
                """INSERT INTO activity_reports VALUES(?,?,?,?,?,?,?,?)
                ON CONFLICT(week) DO UPDATE SET generated_at=excluded.generated_at,
                period_start=excluded.period_start,period_end=excluded.period_end,
                common_path=excluded.common_path,summary_json=excluded.summary_json,
                personal_paths_json=excluded.personal_paths_json""",
                (report.week, report.generated_at, report.period_start, report.period_end,
                 report.common_path, json.dumps(report.summary, ensure_ascii=False),
                 json.dumps(report.personal_paths, ensure_ascii=False), report.sent_at),
            )

    def report(self, week: str | None = None) -> ActivityReport | None:
        with self._connect() as db:
            row = db.execute(
                "SELECT * FROM activity_reports " + ("WHERE week=?" if week else "ORDER BY week DESC LIMIT 1"),
                (week,) if week else (),
            ).fetchone()
        if not row:
            return None
        return ActivityReport(
            row["week"], row["generated_at"], row["period_start"], row["period_end"],
            row["common_path"], json.loads(row["summary_json"]),
            json.loads(row["personal_paths_json"]), row["sent_at"],
        )

    def mark_sent(self, week: str, sent_at: str) -> None:
        with self._connect() as db:
            db.execute("UPDATE activity_reports SET sent_at=? WHERE week=?", (sent_at, week))

    def delivered(self, week: str, user_id: int) -> bool:
        with self._connect() as db:
            return db.execute(
                "SELECT 1 FROM activity_deliveries WHERE week=? AND user_id=?", (week, user_id)
            ).fetchone() is not None

    def mark_delivered(self, week: str, user_id: int, sent_at: str) -> None:
        with self._connect() as db:
            db.execute(
                "INSERT OR IGNORE INTO activity_deliveries VALUES(?,?,?)", (week, user_id, sent_at)
            )


def fetch_changed_shipments(client: MoySkladClient, updated_from: str | None) -> list[dict]:
    filters = [f"moment>={ACTIVITY_START:%Y-%m-%d 00:00:00}"]
    if updated_from:
        filters.append(f"updated>={updated_from}")
    return client._rows("entity/demand", {
        # MoySklad restricts page size for requests with expanded entities.
        # The shared client follows nextHref, so 100 still loads the full set.
        "limit": 100, "filter": ";".join(filters), "expand": "agent,salesChannel",
    })


def sync_shipments(database: ActivityDatabase, client: MoySkladClient, now: datetime) -> int:
    cursor = database.meta("shipments_cursor")
    if database.shipment_count() == 0:
        # Recover automatically if an interrupted/empty first sync left a cursor.
        cursor = None
    updated_from = None
    if cursor:
        try:
            updated_from = (datetime.strptime(cursor[:19], "%Y-%m-%d %H:%M:%S") - timedelta(seconds=5)).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            updated_from = None
    next_cursor = now.strftime("%Y-%m-%d %H:%M:%S")
    count = database.upsert_shipments(fetch_changed_shipments(client, updated_from))
    if count or database.shipment_count():
        database.set_meta("shipments_cursor", next_cursor)
    return count


def _money(minor: int) -> float:
    return float(Decimal(minor) / Decimal(100))


def _snapshot(rows: list[sqlite3.Row], report_monday: date) -> dict:
    period_end = datetime.combine(report_monday, datetime.min.time())
    period_start = period_end - timedelta(days=7)
    grouped: dict[str, list] = defaultdict(list)
    for row in rows:
        moment = datetime.strptime(row["moment"][:19], "%Y-%m-%d %H:%M:%S")
        if moment < period_end:
            grouped[row["client_key"]].append((moment, row))
    new, returned, lost = [], [], []
    for client_rows in grouped.values():
        client_rows.sort(key=lambda item: item[0])
        first_moment, _ = client_rows[0]
        last_moment, last = client_rows[-1]
        week_rows = [item for item in client_rows if period_start <= item[0] < period_end]
        base = {
            "client": last["client_name"], "manager": last["channel_name"],
            "channel_href": last["channel_href"], "counterparties": sorted({item[1]["agent_name"] for item in client_rows}),
        }
        if first_moment >= period_start:
            new.append({**base, "first": first_moment, "count": len(week_rows),
                        "amount": sum(_money(item[1]["amount_minor"]) for item in week_rows)})
        if week_rows:
            first_week = week_rows[0][0]
            previous = [item for item in client_rows if item[0] < period_start]
            if previous and (first_week.date() - previous[-1][0].date()).days >= 20:
                returned.append({**base, "returned": first_week,
                                 "days": (first_week.date() - previous[-1][0].date()).days,
                                 "amount": sum(_money(item[1]["amount_minor"]) for item in week_rows)})
        days = (report_monday - last_moment.date()).days
        if days >= 20:
            category = "20–29 дней" if days < 30 else "30–59 дней" if days < 60 else "60+ дней"
            lost.append({**base, "last": last_moment, "days": days, "category": category,
                         "last_amount": _money(last["amount_minor"])})
    return {"new": new, "returned": returned, "lost": lost}


def report_rows(rows: list[sqlite3.Row], report_monday: date) -> dict:
    current = _snapshot(rows, report_monday)
    previous = _snapshot(rows, report_monday - timedelta(days=7))
    return {
        **current,
        "previous": previous,
        "period_start": report_monday - timedelta(days=7),
        "period_end": report_monday - timedelta(days=1),
        "previous_start": report_monday - timedelta(days=14),
        "previous_end": report_monday - timedelta(days=8),
    }


def _sheet_table(sheet, headers: list[str], rows: list[list], widths: list[int]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="263238")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width


def build_activity_excel(destination: Path, data: dict, channel_href: str = "") -> dict:
    def chosen(items):
        return [item for item in items if not channel_href or _canonical_href(item["channel_href"]) == _canonical_href(channel_href)]
    new, returned, lost = chosen(data["new"]), chosen(data["returned"]), chosen(data["lost"])
    previous_new = chosen(data["previous"]["new"])
    previous_returned = chosen(data["previous"]["returned"])
    previous_lost = chosen(data["previous"]["lost"])
    managers = sorted({item["manager"] or "Без менеджера" for group in (new, returned, lost, previous_new, previous_returned, previous_lost) for item in group})
    wb = Workbook()
    dash = wb.active
    dash.title = "Дэшборд"
    dash.merge_cells("A1:O1")
    dash["A1"] = "Новые кенты и кенты-потеряшки"
    dash["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    dash["A1"].fill = PatternFill("solid", fgColor="263238")
    dash["A2"] = f"Отчёт за {data['period_start']:%d.%m.%Y}–{data['period_end']:%d.%m.%Y}"
    dash.merge_cells("A2:O2")
    headers = ["Менеджер", "Новые кенты — эта неделя", "Новые — предыдущая", "Изменение",
               "Выручка новых", "Вернувшиеся — эта неделя", "Вернувшиеся — предыдущая", "Изменение",
               "Выручка вернувшихся",
               "Клиенты, которые давно не заказывали: 20–29 дней",
               "Клиенты, которые давно не заказывали: 30–59 дней",
               "Клиенты, которые давно не заказывали: 60+ дней",
               "Потеряшки — всего сейчас", "Потеряшки — предыдущая неделя", "Изменение"]
    dash.append(headers)
    for manager in managers:
        n = [x for x in new if (x["manager"] or "Без менеджера") == manager]
        r = [x for x in returned if (x["manager"] or "Без менеджера") == manager]
        l = [x for x in lost if (x["manager"] or "Без менеджера") == manager]
        pn = [x for x in previous_new if (x["manager"] or "Без менеджера") == manager]
        pr = [x for x in previous_returned if (x["manager"] or "Без менеджера") == manager]
        pl = [x for x in previous_lost if (x["manager"] or "Без менеджера") == manager]
        dash.append([manager, len(n), len(pn), len(n) - len(pn), sum(x["amount"] for x in n),
                     len(r), len(pr), len(r) - len(pr), sum(x["amount"] for x in r),
                     sum(x["category"] == "20–29 дней" for x in l), sum(x["category"] == "30–59 дней" for x in l),
                     sum(x["category"] == "60+ дней" for x in l), len(l), len(pl), len(l) - len(pl)])
    dash.append(["ИТОГО", len(new), len(previous_new), len(new) - len(previous_new), sum(x["amount"] for x in new),
                 len(returned), len(previous_returned), len(returned) - len(previous_returned), sum(x["amount"] for x in returned),
                 sum(x["category"] == "20–29 дней" for x in lost), sum(x["category"] == "30–59 дней" for x in lost),
                 sum(x["category"] == "60+ дней" for x in lost), len(lost), len(previous_lost), len(lost) - len(previous_lost)])
    for cell in dash[3]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="263238"); cell.alignment = Alignment(wrap_text=True)
    dash.freeze_panes = "B4"
    for col, width in zip("ABCDEFGHIJKLMNO", [24,20,20,14,20,22,22,14,23,34,34,34,22,26,14]): dash.column_dimensions[col].width = width
    for row in dash.iter_rows(min_row=4):
        row[4].number_format = row[8].number_format = '#,##0.00 [$₽-ru-RU]'
    ws = wb.create_sheet("Новые кенты")
    _sheet_table(ws, ["Клиент", "Менеджер", "Первая отгрузка", "Отгрузок за неделю", "Сумма", "Контрагенты"],
                 [[x["client"], x["manager"], x["first"], x["count"], x["amount"], "\n".join(x["counterparties"])] for x in new], [28,22,20,20,18,55])
    ws = wb.create_sheet("Вернувшиеся кенты")
    _sheet_table(ws, ["Клиент", "Менеджер", "Дата возвращения", "Перерыв, дней", "Сумма за неделю", "Контрагенты"],
                 [[x["client"], x["manager"], x["returned"], x["days"], x["amount"], "\n".join(x["counterparties"])] for x in returned], [28,22,20,18,20,55])
    ws = wb.create_sheet("Кенты-потеряшки")
    _sheet_table(ws, ["Клиент", "Менеджер", "Категория", "Последняя отгрузка", "Дней без отгрузок", "Сумма последней отгрузки", "Контрагенты"],
                 [[x["client"], x["manager"], x["category"], x["last"], x["days"], x["last_amount"], "\n".join(x["counterparties"])] for x in lost], [28,22,18,22,22,25,55])
    colors = {"20–29 дней": "FFF2CC", "30–59 дней": "FCE4D6", "60+ дней": "F4CCCC"}
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.fill = PatternFill("solid", fgColor=colors.get(row[2].value, "FFFFFF"))
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination)
    return {
        "new": len(new), "previous_new": len(previous_new),
        "returned": len(returned), "previous_returned": len(previous_returned),
        "lost": len(lost), "previous_lost": len(previous_lost),
    }


def build_weekly_report(
    database: ActivityDatabase,
    client: MoySkladClient,
    storage: Path,
    monday: date,
    manager_channels: dict[str, str],
    now: datetime,
    *,
    save: bool = True,
) -> ActivityReport:
    synced = sync_shipments(database, client, now)
    data = report_rows(database.shipments(), monday)
    folder = storage / monday.isoformat()
    common_path = folder / f"Кенты_{monday:%d.%m.%Y}_общий.xlsx"
    summary = build_activity_excel(common_path, data)
    personal_paths = {}
    personal_summary = {}
    for access_id, channel_href in manager_channels.items():
        path = folder / f"Кенты_{monday:%d.%m.%Y}_менеджер_{access_id}.xlsx"
        personal_summary[access_id] = build_activity_excel(path, data, channel_href)
        personal_paths[access_id] = str(path)
    summary["personal"] = personal_summary
    report = ActivityReport(
        week=monday.isoformat(), generated_at=now.isoformat(timespec="seconds"),
        period_start=data["period_start"].isoformat(), period_end=data["period_end"].isoformat(),
        common_path=str(common_path), summary=summary, personal_paths=personal_paths, sent_at=None,
    )
    report.summary["synced"] = synced
    report.summary["stored"] = len(database.shipments())
    if save:
        database.save_report(report)
    return report
