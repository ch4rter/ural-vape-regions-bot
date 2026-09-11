"""Read-only MoySklad bonus reports and monthly folder classification."""

from __future__ import annotations

import json
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from bonus_report_v2 import build_bonus_workbook_v2

from moysklad_price import BONUS_CATEGORIES, MoySkladClient, MoySkladError, _canonical_href


OTHER_CATEGORY = "Прочее"
UNCLASSIFIED_CATEGORY = "Не классифицировано"
EXCLUDED_CATEGORY = "Не учитывать"
ALL_CHANNELS = "Все менеджеры"
MAX_BONUS_CATEGORIES = 20


@dataclass(frozen=True)
class BonusClassification:
    categories: tuple[str, ...]
    folders: dict[str, str]


@dataclass(frozen=True)
class SalesChannelOption:
    name: str
    href: str


@dataclass(frozen=True)
class BonusDocument:
    kind: str
    moment: str
    name: str
    agent: str
    state: str
    sales_channel: str
    total: Decimal
    paid: Decimal
    categories: dict[str, Decimal]


@dataclass(frozen=True)
class BonusReportResult:
    path: Path
    document_count: int
    shipment_count: int
    return_count: int
    unclassified_paths: tuple[str, ...]


def month_bounds(month: str) -> tuple[str, str]:
    try:
        year, number = (int(part) for part in month.split("-", 1))
        start = datetime(year, number, 1)
    except (TypeError, ValueError) as error:
        raise MoySkladError("Месяц должен быть указан в формате ГГГГ-ММ.") from error
    if number == 12:
        end = datetime(year + 1, 1, 1)
    else:
        end = datetime(year, number + 1, 1)
    return start.strftime("%Y-%m-%d 00:00:00"), end.strftime("%Y-%m-%d 00:00:00")


def month_title(month: str) -> str:
    year, number = (int(part) for part in month.split("-", 1))
    names = (
        "январь", "февраль", "март", "апрель", "май", "июнь",
        "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь",
    )
    return f"{names[number - 1]} {year}"


def _full_folder_path(folder: dict) -> str:
    name = str(folder.get("name", "")).strip()
    parent = str(folder.get("pathName", "")).strip(" /\t")
    return f"{parent}/{name}" if parent else name


def leaf_folders(folders: list[dict]) -> list[dict]:
    active = [folder for folder in folders if not folder.get("archived") and folder.get("name")]
    paths = {_full_folder_path(folder).casefold() for folder in active}
    result = []
    for folder in active:
        path = _full_folder_path(folder)
        prefix = path.casefold() + "/"
        if not any(other.startswith(prefix) for other in paths):
            result.append(folder)
    return sorted(result, key=lambda folder: _full_folder_path(folder).casefold())


def _validate_categories(values: list[str]) -> tuple[str, ...]:
    categories = []
    seen = set()
    for value in values:
        category = str(value or "").strip()
        if not category:
            continue
        key = category.casefold()
        if key == EXCLUDED_CATEGORY.casefold() and category != EXCLUDED_CATEGORY:
            raise MoySkladError(f"Служебную категорию нужно оставить точно как «{EXCLUDED_CATEGORY}».")
        if key in seen:
            raise MoySkladError(f"В справочнике категория «{category}» указана повторно.")
        if len(category) > 60:
            raise MoySkladError(f"Название категории «{category}» длиннее 60 символов.")
        seen.add(key)
        categories.append(category)
    if EXCLUDED_CATEGORY.casefold() not in seen:
        raise MoySkladError(f"В справочнике должна оставаться служебная категория «{EXCLUDED_CATEGORY}».")
    if len(categories) < 2:
        raise MoySkladError("Добавьте в справочник хотя бы одну премиальную категорию.")
    if len(categories) > MAX_BONUS_CATEGORIES + 1:
        raise MoySkladError(f"Допускается не более {MAX_BONUS_CATEGORIES} премиальных категорий.")
    return tuple(categories)


def read_classification(path: Path) -> BonusClassification:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    if len(workbook.worksheets) > 1:
        category_values = [
            str(row[0].value or "").strip()
            for row in workbook.worksheets[1].iter_rows(
                min_row=2, min_col=1, max_col=1
            )
        ]
        categories = _validate_categories(category_values)
    else:
        categories = tuple(BONUS_CATEGORIES)
    headers = {
        str(cell.value).strip().casefold(): index
        for index, cell in enumerate(sheet[1])
        if cell.value is not None
    }
    path_index = headers.get("полный путь папки")
    category_index = headers.get("категория")
    if path_index is None or category_index is None:
        workbook.close()
        raise MoySkladError("В таблице нужны колонки «Полный путь папки» и «Категория».")
    entered: dict[str, str] = {}
    errors = []
    allowed = set(categories)
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        folder_path = str(values[path_index] or "").strip(" /\t")
        category = str(values[category_index] or "").strip()
        if not folder_path:
            continue
        if not category:
            errors.append(f"строка {row_number}: не выбрана категория")
            continue
        if category not in allowed:
            errors.append(f"строка {row_number}: неизвестная категория «{category}»")
            continue
        key = folder_path.casefold()
        if key in entered and entered[key] != category:
            errors.append(f"строка {row_number}: папка указана повторно с другой категорией")
        entered[key] = category
    workbook.close()
    if errors:
        details = "; ".join(errors[:8])
        if len(errors) > 8:
            details += f"; и ещё {len(errors) - 8}"
        raise MoySkladError("Таблица классификации не прошла проверку: " + details + ".")
    if not entered:
        raise MoySkladError("В таблице нет заполненных товарных папок.")
    # Older exports contained the whole tree. Keep only terminal folders so the
    # saved monthly rules have the same semantics as new leaf-only templates.
    mapping = {
        key: category
        for key, category in entered.items()
        if not any(other.startswith(key + "/") for other in entered)
    }
    return BonusClassification(categories=categories, folders=mapping)


def save_classification(source: Path, destination: Path) -> int:
    classification = read_classification(source)
    destination.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "version": 2,
        "categories": list(classification.categories),
        "folders": classification.folders,
    }
    temporary = destination.with_suffix(".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(destination)
    return len(classification.folders)


def load_classification(path: Path) -> BonusClassification:
    if not path.exists():
        raise MoySkladError("Для выбранного месяца ещё не загружена классификация товарных папок.")
    payload = json.loads(path.read_text(encoding="utf-8"))
    folders = payload.get("folders", {})
    if not isinstance(folders, dict) or not folders:
        raise MoySkladError("Сохранённая классификация товарных папок повреждена или пуста.")
    raw_categories = payload.get("categories")
    if isinstance(raw_categories, list):
        categories = _validate_categories([str(value) for value in raw_categories])
    else:
        # Version 1 stored only folder mappings and used the original fixed list.
        categories = _validate_categories(list(BONUS_CATEGORIES))
    normalized = {str(key).casefold(): str(value) for key, value in folders.items()}
    unknown = sorted(set(normalized.values()) - set(categories))
    if unknown:
        raise MoySkladError("В сохранённой классификации неизвестные категории: " + ", ".join(unknown) + ".")
    return BonusClassification(categories=categories, folders=normalized)


def _expanded_rows(
    client: MoySkladClient,
    endpoint: str,
    start: str,
    end: str,
    expand: str,
    sales_channel_href: str = "",
) -> list[dict]:
    filters = f"moment>={start};moment<{end};applicable=true"
    if sales_channel_href:
        filters += f";salesChannel={sales_channel_href}"
    return client._rows(
        endpoint,
        {
            "limit": 100,
            "filter": filters,
            "expand": expand,
        },
    )


def fetch_month_documents(
    client: MoySkladClient,
    month: str,
    sales_channel_href: str = "",
) -> tuple[list[dict], list[dict]]:
    start, end = month_bounds(month)
    # Positions are intentionally not expanded here. This first, lightweight
    # request is used to choose a sales channel. Position rows are fetched only
    # for documents included in the requested report.
    try:
        demands = _expanded_rows(
            client,
            "entity/demand",
            start,
            end,
            "agent,state,salesChannel",
            sales_channel_href,
        )
    except MoySkladError as error:
        # Some account revisions may not accept filtering demands by a sales
        # channel. Fall back to lightweight headers and filter locally; product
        # positions are still fetched only for the selected manager.
        if not sales_channel_href or error.status != 400:
            raise
        demands = _expanded_rows(
            client, "entity/demand", start, end, "agent,state,salesChannel"
        )
    # The API user may intentionally have no permission to view customer
    # returns. Bonus reports therefore operate on shipments only.
    return demands, []


def available_sales_channels(client: MoySkladClient) -> tuple[SalesChannelOption, ...]:
    options = {
        (
            str(row.get("name", "")).strip(),
            _canonical_href(str(row.get("meta", {}).get("href", ""))),
        )
        for row in client.sales_channels()
        if not row.get("archived")
        and str(row.get("name", "")).strip()
        and str(row.get("meta", {}).get("href", "")).strip()
    }
    return tuple(
        SalesChannelOption(name=name, href=href)
        for name, href in sorted(options, key=lambda row: row[0].casefold())
    )


def _name(value: object, default: str = "—") -> str:
    return str(value.get("name", default)).strip() if isinstance(value, dict) else default


def _channel(document: dict) -> str:
    channel = _name(document.get("salesChannel"), "")
    if channel:
        return channel
    demand = document.get("demand")
    return _name(demand.get("salesChannel"), "") if isinstance(demand, dict) else ""


def _channel_href(document: dict) -> str:
    channel = document.get("salesChannel")
    if isinstance(channel, dict):
        return _canonical_href(str(channel.get("meta", {}).get("href", "")))
    return ""


def sales_channels(documents: tuple[list[dict], list[dict]]) -> tuple[str, ...]:
    values = {_channel(document) for rows in documents for document in rows}
    return tuple(sorted((value for value in values if value), key=str.casefold))


def _minor_money(value: object) -> Decimal:
    try:
        return Decimal(str(value or 0)) / Decimal(100)
    except Exception:
        return Decimal(0)


def _position_sum(position: dict) -> Decimal:
    quantity = Decimal(str(position.get("quantity") or 0))
    price = _minor_money(position.get("price"))
    discount = Decimal(str(position.get("discount") or 0))
    return price * quantity * (Decimal(1) - discount / Decimal(100))


def _assortment_paths(client: MoySkladClient) -> dict[str, str]:
    result = {}
    for item in client.assortment():
        href = _canonical_href(str(item.get("meta", {}).get("href", "")))
        path = str(item.get("pathName", "")).strip(" /\t")
        if href:
            result[href] = path
    return result


def _document_positions(document: dict, client: MoySkladClient) -> list[dict]:
    positions = document.get("positions", {})
    if isinstance(positions, dict) and isinstance(positions.get("rows"), list):
        rows = positions["rows"]
        meta = positions.get("meta", {})
        if int(meta.get("size", len(rows)) or 0) <= len(rows):
            return [row for row in rows if isinstance(row, dict)]
        href = meta.get("href")
    else:
        href = positions.get("meta", {}).get("href") if isinstance(positions, dict) else None
    if not href:
        return []
    payload = client._get_json(str(href), {"limit": 1000})
    return [row for row in payload.get("rows", []) if isinstance(row, dict)]


def prepare_documents(
    client: MoySkladClient,
    raw: tuple[list[dict], list[dict]],
    classification: BonusClassification,
    channel: str,
    channel_href: str = "",
) -> tuple[list[BonusDocument], set[str]]:
    assortment_paths = _assortment_paths(client)
    result: list[BonusDocument] = []
    unknown: set[str] = set()
    report_categories = tuple(
        category for category in classification.categories
        if category != EXCLUDED_CATEGORY
    )
    selected = []
    for document in raw[0]:
        document_channel = _channel(document)
        matches = (
            _channel_href(document) == _canonical_href(channel_href)
            if channel_href else document_channel.casefold() == channel.casefold()
        )
        if channel == ALL_CHANNELS or matches:
            selected.append(document)
    # Position endpoints are independent. A small bounded pool considerably
    # reduces report time without creating an aggressive burst against the API.
    with ThreadPoolExecutor(max_workers=min(4, max(1, len(selected)))) as executor:
        position_sets = list(executor.map(
            lambda document: _document_positions(document, client), selected
        ))
    for document, positions in zip(selected, position_sets):
        document_channel = _channel(document)
        amounts = {category: Decimal(0) for category in report_categories}
        amounts[OTHER_CATEGORY] = Decimal(0)
        amounts[UNCLASSIFIED_CATEGORY] = Decimal(0)
        for position in positions:
            assortment = position.get("assortment", {})
            href = _canonical_href(str(assortment.get("meta", {}).get("href", "")))
            folder_path = assortment_paths.get(href, "")
            category = classification.folders.get(folder_path.casefold()) if folder_path else None
            value = _position_sum(position)
            if category in amounts:
                amounts[category] += value
            elif category == EXCLUDED_CATEGORY:
                amounts[OTHER_CATEGORY] += value
            else:
                amounts[UNCLASSIFIED_CATEGORY] += value
                unknown.add(folder_path or "Без папки")
        result.append(BonusDocument(
            kind="Отгрузка",
            moment=str(document.get("moment", "")),
            name=str(document.get("name", "")),
            agent=_name(document.get("agent")),
            state=_name(document.get("state")),
            sales_channel=document_channel,
            total=_minor_money(document.get("sum")),
            paid=_minor_money(document.get("payedSum")),
            categories=amounts,
        ))
    result.sort(key=lambda row: (row.moment, row.kind, row.name))
    return result, unknown


def _detail_headers(category_headers: list[str]) -> list[str]:
    return [
        "Тип", "Дата", "Документ", "Контрагент", "Статус", "Канал продаж",
        "Сумма документа", "Оплачено", *category_headers, OTHER_CATEGORY,
        UNCLASSIFIED_CATEGORY, "Контрольное расхождение",
    ]


def _fill_detail_sheet(
    sheet,
    documents: list[BonusDocument],
    category_headers: list[str],
    title: str,
    month: str,
) -> None:
    headers = _detail_headers(category_headers)
    sheet.append([title, month_title(month)])
    sheet.append([])
    sheet.append(headers)
    blue = PatternFill("solid", fgColor="2F5597")
    total_fill = PatternFill("solid", fgColor="D9EAD3")
    for cell in sheet[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = blue
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for document in documents:
        categorized = sum(document.categories.values(), Decimal(0))
        sheet.append([
            document.kind, document.moment[:10], document.name, document.agent,
            document.state, document.sales_channel, float(document.total),
            float(document.paid),
            *(float(document.categories[category]) for category in category_headers),
            float(document.categories[OTHER_CATEGORY]),
            float(document.categories[UNCLASSIFIED_CATEGORY]),
            float(document.total - categorized),
        ])
    first_data_row = 4
    total_row = sheet.max_row + 1
    sheet.cell(total_row, 1, "ИТОГО")
    for column in range(7, len(headers) + 1):
        letter = get_column_letter(column)
        sheet.cell(total_row, column, f"=SUM({letter}{first_data_row}:{letter}{total_row - 1})")
    for cell in sheet[total_row]:
        cell.font = Font(bold=True)
        cell.fill = total_fill
    for row in range(first_data_row, total_row + 1):
        for column in range(7, len(headers) + 1):
            sheet.cell(row, column).number_format = '#,##0.00'
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:{sheet.cell(3, len(headers)).column_letter}{total_row - 1}"
    widths = [13, 13, 18, 38, 22, 24, 19, 16] + [24] * len(category_headers) + [15, 23, 23]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width


def _safe_sheet_title(value: str, used: set[str]) -> str:
    cleaned = "".join("_" if char in "[]:*?/\\" else char for char in value).strip() or "Без канала"
    base = cleaned[:31]
    candidate = base
    number = 2
    while candidate.casefold() in used:
        suffix = f" {number}"
        candidate = base[:31 - len(suffix)] + suffix
        number += 1
    used.add(candidate.casefold())
    return candidate


def _fill_summary_sheet(
    sheet,
    documents: list[BonusDocument],
    category_headers: list[str],
    month: str,
    channel: str,
) -> None:
    dark_blue = PatternFill("solid", fgColor="1F4E78")
    blue = PatternFill("solid", fgColor="5B9BD5")
    light_blue = PatternFill("solid", fgColor="DDEBF7")
    green = PatternFill("solid", fgColor="E2F0D9")
    white_font = Font(bold=True, color="FFFFFF")
    thin_gray = Side(style="thin", color="D9E1F2")
    table_border = Border(bottom=thin_gray)
    money_format = '#,##0.00'

    total = sum((row.total for row in documents), Decimal(0))
    paid = sum((row.paid for row in documents), Decimal(0))
    agents = {row.agent.casefold() for row in documents if row.agent and row.agent != "—"}
    average = total / len(documents) if documents else Decimal(0)
    report_name = "Все менеджеры" if channel == ALL_CHANNELS else channel

    last_column = max(8, 7 + len(category_headers))
    last_letter = sheet.cell(1, last_column).column_letter
    sheet.merge_cells(f"A1:{last_letter}1")
    sheet["A1"] = "ПРЕМИАЛЬНЫЙ ОТЧЁТ"
    sheet["A1"].font = Font(bold=True, color="FFFFFF", size=16)
    sheet["A1"].fill = dark_blue
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet["A2"] = "Период"
    sheet["B2"] = month_title(month).capitalize()
    sheet["D2"] = "Канал продаж"
    sheet["E2"] = report_name
    for coordinate in ("A2", "D2"):
        sheet[coordinate].font = Font(bold=True, color="1F4E78")

    metrics = [
        ("Отгрузок", len(documents), "0"),
        ("Контрагентов", len(agents), "0"),
        ("Сумма отгрузок", float(total), money_format),
        ("Оплачено", float(paid), money_format),
        ("Задолженность", float(total - paid), money_format),
        ("Средний чек", float(average), money_format),
    ]
    for index, (label, value, number_format) in enumerate(metrics, 1):
        cell = sheet.cell(4, index)
        cell.value = label
        cell.font = white_font
        cell.fill = blue
        cell.alignment = Alignment(horizontal="center")
        value_cell = sheet.cell(5, index)
        value_cell.value = value
        value_cell.font = Font(bold=True, color="1F4E78", size=12)
        value_cell.fill = light_blue
        value_cell.alignment = Alignment(horizontal="center")
        value_cell.number_format = number_format

    manager_headers = [
        "Менеджер", "Контрагентов", "Отгрузок", "Сумма", "Оплачено", "Задолженность",
        "Средний чек", *category_headers, OTHER_CATEGORY, UNCLASSIFIED_CATEGORY,
    ]
    manager_start = 8
    sheet.cell(manager_start, 1, "Сводка по менеджерам")
    sheet.cell(manager_start, 1).font = Font(bold=True, color="1F4E78", size=12)
    manager_header_row = manager_start + 1
    for column, header in enumerate(manager_headers, 1):
        cell = sheet.cell(manager_header_row, column, header)
        cell.font = white_font
        cell.fill = dark_blue
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    grouped: dict[str, list[BonusDocument]] = {}
    for document in documents:
        grouped.setdefault(document.sales_channel or "Без канала", []).append(document)
    manager_first_row = manager_header_row + 1
    for manager in sorted(grouped, key=str.casefold):
        rows = grouped[manager]
        manager_total = sum((row.total for row in rows), Decimal(0))
        manager_paid = sum((row.paid for row in rows), Decimal(0))
        category_totals = [
            sum((row.categories[category] for row in rows), Decimal(0))
            for category in category_headers
        ]
        sheet.append([
            manager,
            len({row.agent.casefold() for row in rows if row.agent and row.agent != "—"}),
            len(rows), float(manager_total), float(manager_paid),
            float(manager_total - manager_paid),
            float(manager_total / len(rows) if rows else 0),
            *(float(value) for value in category_totals),
            float(sum((row.categories[OTHER_CATEGORY] for row in rows), Decimal(0))),
            float(sum((row.categories[UNCLASSIFIED_CATEGORY] for row in rows), Decimal(0))),
        ])
    manager_total_row = sheet.max_row + 1
    sheet.cell(manager_total_row, 1, "ИТОГО")
    sheet.cell(manager_total_row, 2, len(agents))
    for column in range(3, len(manager_headers) + 1):
        letter = get_column_letter(column)
        sheet.cell(manager_total_row, column, f"=SUM({letter}{manager_first_row}:{letter}{manager_total_row - 1})")
    sheet.cell(manager_total_row, 7, f"=IFERROR(D{manager_total_row}/C{manager_total_row},0)")
    for cell in sheet[manager_total_row]:
        cell.font = Font(bold=True)
        cell.fill = green

    category_start = manager_total_row + 3
    sheet.cell(category_start, 1, "Продажи по категориям")
    sheet.cell(category_start, 1).font = Font(bold=True, color="1F4E78", size=12)
    category_header_row = category_start + 1
    for column, header in enumerate(("Категория", "Сумма", "Доля продаж", "Контрагентов"), 1):
        cell = sheet.cell(category_header_row, column, header)
        cell.font = white_font
        cell.fill = dark_blue
        cell.alignment = Alignment(horizontal="center")
    all_categories = [*category_headers, OTHER_CATEGORY, UNCLASSIFIED_CATEGORY]
    category_first_row = category_header_row + 1
    for category in all_categories:
        amount = sum((row.categories[category] for row in documents), Decimal(0))
        buyers = {
            (row.sales_channel.casefold(), row.agent.casefold())
            for row in documents
            if row.categories[category] != 0
        }
        sheet.append([
            category, float(amount), float(amount / total if total else 0), len(buyers)
        ])

    counterparty_start = sheet.max_row + 3
    sheet.cell(counterparty_start, 1, "Статистика по контрагентам")
    sheet.cell(counterparty_start, 1).font = Font(bold=True, color="1F4E78", size=12)
    include_manager = channel == ALL_CHANNELS
    counterparty_headers = ["Контрагент"]
    if include_manager:
        counterparty_headers.append("Менеджер")
    counterparty_headers.extend([
        "Отгрузок", "Сумма", "Оплачено", "Задолженность", "Средний чек",
        *category_headers, OTHER_CATEGORY, UNCLASSIFIED_CATEGORY,
    ])
    counterparty_header_row = counterparty_start + 1
    for column, header in enumerate(counterparty_headers, 1):
        cell = sheet.cell(counterparty_header_row, column, header)
        cell.font = white_font
        cell.fill = dark_blue
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    counterparties: dict[tuple[str, str], list[BonusDocument]] = {}
    for document in documents:
        key = (document.sales_channel.casefold() if include_manager else "", document.agent.casefold())
        counterparties.setdefault(key, []).append(document)
    ordered_counterparties = sorted(
        counterparties.values(),
        key=lambda rows: (-sum((row.total for row in rows), Decimal(0)), rows[0].agent.casefold()),
    )
    counterparty_first_row = counterparty_header_row + 1
    for rows in ordered_counterparties:
        client_total = sum((row.total for row in rows), Decimal(0))
        client_paid = sum((row.paid for row in rows), Decimal(0))
        values: list[object] = [rows[0].agent]
        if include_manager:
            values.append(rows[0].sales_channel or "Без канала")
        values.extend([
            len(rows), float(client_total), float(client_paid), float(client_total - client_paid),
            float(client_total / len(rows) if rows else 0),
            *(float(sum((row.categories[category] for row in rows), Decimal(0))) for category in category_headers),
            float(sum((row.categories[OTHER_CATEGORY] for row in rows), Decimal(0))),
            float(sum((row.categories[UNCLASSIFIED_CATEGORY] for row in rows), Decimal(0))),
        ])
        sheet.append(values)
    counterparty_total_row = sheet.max_row + 1
    sheet.cell(counterparty_total_row, 1, "ИТОГО")
    orders_column = 3 if include_manager else 2
    for column in range(orders_column, len(counterparty_headers) + 1):
        letter = get_column_letter(column)
        sheet.cell(counterparty_total_row, column, f"=SUM({letter}{counterparty_first_row}:{letter}{counterparty_total_row - 1})")
    average_column = 7 if include_manager else 6
    sum_column = 4 if include_manager else 3
    sheet.cell(
        counterparty_total_row, average_column,
        f"=IFERROR({sheet.cell(counterparty_total_row, sum_column).coordinate}/{sheet.cell(counterparty_total_row, orders_column).coordinate},0)",
    )
    for cell in sheet[counterparty_total_row]:
        cell.font = Font(bold=True)
        cell.fill = green

    money_columns_manager = range(4, len(manager_headers) + 1)
    for row in range(manager_first_row, manager_total_row + 1):
        for column in money_columns_manager:
            sheet.cell(row, column).number_format = money_format
            sheet.cell(row, column).border = table_border
    for row in range(category_first_row, category_first_row + len(all_categories)):
        sheet.cell(row, 2).number_format = money_format
        sheet.cell(row, 3).number_format = "0.0%"
    for row in range(counterparty_first_row, counterparty_total_row + 1):
        for column in range(sum_column, len(counterparty_headers) + 1):
            sheet.cell(row, column).number_format = money_format
            sheet.cell(row, column).border = table_border

    if all_categories:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Продажи по категориям"
        chart.y_axis.title = "Категория"
        chart.x_axis.title = "Сумма"
        chart.height = 7
        chart.width = 13
        data = Reference(sheet, min_col=2, min_row=category_header_row, max_row=category_first_row + len(all_categories) - 1)
        labels = Reference(sheet, min_col=1, min_row=category_first_row, max_row=category_first_row + len(all_categories) - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(labels)
        sheet.add_chart(chart, "F12")

    sheet.freeze_panes = f"A{counterparty_first_row}"
    sheet.auto_filter.ref = (
        f"A{counterparty_header_row}:"
        f"{sheet.cell(counterparty_header_row, len(counterparty_headers)).column_letter}{counterparty_total_row - 1}"
    )
    widths = [38, 25, 13, 18, 18, 18, 18] + [22] * (len(category_headers) + 2)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.sheet_view.showGridLines = False


def build_bonus_report_legacy(
    token: str,
    month: str,
    channel: str,
    classification_path: Path,
    destination: Path,
    *,
    sales_channel_href: str = "",
    client: MoySkladClient | None = None,
    raw_documents: tuple[list[dict], list[dict]] | None = None,
) -> BonusReportResult:
    client = client or MoySkladClient(token)
    classification = load_classification(classification_path)
    raw = raw_documents or fetch_month_documents(client, month, sales_channel_href)
    documents, unknown = prepare_documents(
        client, raw, classification, channel, sales_channel_href
    )
    if not documents:
        raise MoySkladError(f"За {month_title(month)} по каналу продаж «{channel}» проведённых документов не найдено.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Отчёт"
    category_headers = [
        category for category in classification.categories
        if category != EXCLUDED_CATEGORY
    ]
    _fill_detail_sheet(
        sheet, documents, category_headers,
        f"Отчёт по премиальным категориям · {channel}", month,
    )

    summary = workbook.create_sheet("Сводка", 0)
    _fill_summary_sheet(summary, documents, category_headers, month, channel)

    if channel == ALL_CHANNELS:
        used_titles = {item.title.casefold() for item in workbook.worksheets}
        grouped: dict[str, list[BonusDocument]] = {}
        for document in documents:
            grouped.setdefault(document.sales_channel or "Без канала", []).append(document)
        for manager in sorted(grouped, key=str.casefold):
            manager_sheet = workbook.create_sheet(_safe_sheet_title(manager, used_titles))
            _fill_detail_sheet(
                manager_sheet, grouped[manager], category_headers,
                f"Отгрузки · {manager}", month,
            )

    if unknown:
        unknown_sheet = workbook.create_sheet("Не классифицировано")
        unknown_sheet.append(["Папка товара"])
        for folder_path in sorted(unknown, key=str.casefold):
            unknown_sheet.append([folder_path])
        unknown_sheet.column_dimensions["A"].width = 90
        unknown_sheet.freeze_panes = "A2"
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()
    return BonusReportResult(
        path=destination,
        document_count=len(documents),
        shipment_count=sum(row.kind == "Отгрузка" for row in documents),
        return_count=sum(row.kind == "Возврат" for row in documents),
        unclassified_paths=tuple(sorted(unknown, key=str.casefold)),
    )


def build_bonus_report(
    token: str,
    month: str,
    channel: str,
    classification_path: Path,
    destination: Path,
    *,
    sales_channel_href: str = "",
    client: MoySkladClient | None = None,
    raw_documents: tuple[list[dict], list[dict]] | None = None,
) -> BonusReportResult:
    """Build the production premium report using the established calculation pipeline."""
    client = client or MoySkladClient(token)
    classification = load_classification(classification_path)
    raw = raw_documents or fetch_month_documents(client, month, sales_channel_href)
    documents, unknown = prepare_documents(
        client, raw, classification, channel, sales_channel_href
    )
    if not documents:
        raise MoySkladError(
            f"За {month_title(month)} по каналу продаж «{channel}» проведённых документов не найдено."
        )
    categories = [
        category for category in classification.categories
        if category != EXCLUDED_CATEGORY
    ]
    categories.extend([OTHER_CATEGORY, UNCLASSIFIED_CATEGORY])
    build_bonus_workbook_v2(
        destination,
        documents,
        categories,
        month_title(month).capitalize(),
        "Все менеджеры" if channel == ALL_CHANNELS else channel,
    )
    return BonusReportResult(
        path=destination,
        document_count=len(documents),
        shipment_count=sum(row.kind == "Отгрузка" for row in documents),
        return_count=sum(row.kind == "Возврат" for row in documents),
        unclassified_paths=tuple(sorted(unknown, key=str.casefold)),
    )
