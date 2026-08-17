import json
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from moysklad_bonus import (
    build_bonus_report,
    leaf_folders,
    read_classification,
    save_classification,
)


class FakeBonusClient:
    def assortment(self):
        return [
            {
                "meta": {"href": "https://api.moysklad.ru/api/remap/1.2/entity/product/oggo"},
                "pathName": "ЭС/Жидкости/OGGO",
            },
            {
                "meta": {"href": "https://api.moysklad.ru/api/remap/1.2/entity/product/other"},
                "pathName": "Услуги/Прочее",
            },
        ]


def classification_xlsx(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Полный путь папки", "Категория", "Комментарий", "ID папки МоегоСклада"])
    sheet.append(["ЭС", "Не учитывать", "", "root"])
    sheet.append(["ЭС/Жидкости/OGGO", "OGGO Аромы/жижи", "", "oggo"])
    sheet.append(["Услуги/Прочее", "Не учитывать", "", "other"])
    workbook.save(path)
    workbook.close()


def test_leaf_folders_excludes_parents_and_archived():
    folders = [
        {"name": "ЭС", "pathName": ""},
        {"name": "Жидкости", "pathName": "ЭС"},
        {"name": "Старое", "pathName": "", "archived": True},
    ]
    assert [row["name"] for row in leaf_folders(folders)] == ["Жидкости"]


def test_old_full_tree_classification_is_normalized_to_leaves(tmp_path):
    source = tmp_path / "mapping.xlsx"
    classification_xlsx(source)
    mapping = read_classification(source)
    assert "эс" not in mapping.folders
    assert mapping.folders["эс/жидкости/oggo"] == "OGGO Аромы/жижи"
    destination = tmp_path / "2026-08.json"
    assert save_classification(source, destination) == 2
    assert len(json.loads(destination.read_text(encoding="utf-8"))["folders"]) == 2


def test_custom_month_categories_are_saved_and_used(tmp_path):
    source = tmp_path / "custom.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Полный путь папки", "Категория"])
    sheet.append(["ЭС/Жидкости/OGGO", "Новинки месяца"])
    guide = workbook.create_sheet("Справочник")
    guide.append(["Допустимые категории"])
    guide.append(["Новинки месяца"])
    guide.append(["Не учитывать"])
    workbook.save(source)
    workbook.close()
    destination = tmp_path / "2026-09.json"
    save_classification(source, destination)
    payload = json.loads(destination.read_text(encoding="utf-8"))
    assert payload["categories"] == ["Новинки месяца", "Не учитывать"]
    assert payload["folders"]["эс/жидкости/oggo"] == "Новинки месяца"
    demand = {
        "moment": "2026-09-10 10:00:00",
        "name": "0002",
        "agent": {"name": "Клиент"},
        "state": {"name": "Отгружено"},
        "salesChannel": {"name": "Валера"},
        "sum": 50000,
        "payedSum": 50000,
        "positions": {"rows": [{
            "assortment": {"meta": {
                "href": "https://api.moysklad.ru/api/remap/1.2/entity/product/oggo"
            }},
            "quantity": 1,
            "price": 50000,
        }]},
    }
    report = tmp_path / "custom-report.xlsx"
    build_bonus_report(
        "token", "2026-09", "Валера", destination, report,
        client=FakeBonusClient(), raw_documents=([demand], []),
    )
    workbook = load_workbook(report, read_only=True)
    headers = [cell.value for cell in workbook["Отчёт"][3]]
    assert "Новинки месяца" in headers
    assert "OGGO Аромы/жижи" not in headers
    workbook.close()


def test_build_bonus_report_splits_categories_and_returns(tmp_path):
    source = tmp_path / "mapping.xlsx"
    classification_xlsx(source)
    classification = tmp_path / "2026-08.json"
    save_classification(source, classification)
    assortment = lambda item: {
        "meta": {"href": f"https://api.moysklad.ru/api/remap/1.2/entity/product/{item}"}
    }
    demand = {
        "moment": "2026-08-10 10:00:00",
        "name": "0001",
        "agent": {"name": "Клиент"},
        "state": {"name": "Отгружено"},
        "salesChannel": {"name": "Валера"},
        "sum": 150000,
        "payedSum": 100000,
        "positions": {"rows": [
            {"assortment": assortment("oggo"), "quantity": 2, "price": 50000, "discount": 10},
            {"assortment": assortment("other"), "quantity": 1, "price": 60000, "discount": 0},
        ]},
    }
    returned = {
        "moment": "2026-08-11 10:00:00",
        "name": "В-1",
        "agent": {"name": "Клиент"},
        "state": {"name": "Проведено"},
        "demand": {"salesChannel": {"name": "Валера"}},
        "sum": 50000,
        "payedSum": 50000,
        "positions": {"rows": [
            {"assortment": assortment("oggo"), "quantity": 1, "price": 50000, "discount": 0},
        ]},
    }
    destination = tmp_path / "report.xlsx"
    result = build_bonus_report(
        "token", "2026-08", "Валера", classification, destination,
        client=FakeBonusClient(), raw_documents=([demand], [returned]),
    )
    assert result.shipment_count == 1
    assert result.return_count == 1
    assert result.unclassified_paths == ()
    workbook = load_workbook(destination, data_only=False)
    sheet = workbook["Отчёт"]
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert Decimal(str(sheet.cell(4, headers["OGGO Аромы/жижи"]).value)) == Decimal("900")
    assert Decimal(str(sheet.cell(4, headers["Прочее"]).value)) == Decimal("600")
    assert Decimal(str(sheet.cell(5, headers["OGGO Аромы/жижи"]).value)) == Decimal("-500")
    workbook.close()
