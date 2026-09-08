from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from contracts import (
    ContractsDB,
    accountant_message,
    buyer_intro,
    create_contract_file,
)


def payload(**updates):
    result = {
        "contract_number": "291",
        "contract_date": "08.09.2026",
        "buyer_type": "ip",
        "buyer_name": "Сафин Рафис Рауфович",
        "representative_genitive": "Сафина Рафиса Рауфовича",
        "buyer_inn": "027808041941",
        "buyer_ogrn": "322028000042041",
        "buyer_address": "г. Уфа, ул. Примерная, д. 1",
        "bank_account": "40802810001500362023",
        "correspondent_account": "30101810745374525104",
        "bik": "044525104",
        "bank_name": "ООО «Банк Точка»",
        "edo_id": "2BE66de2b940a0f45bfbe36676dd5ed4d9a",
    }
    result.update(updates)
    return result


def document_text(path: Path) -> str:
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def test_contract_template_is_filled_and_keeps_standard_body(tmp_path):
    root = Path(__file__).resolve().parents[1]
    destination = create_contract_file(
        root / "templates/contract_seletkov.docx", tmp_path, payload()
    )
    text = document_text(destination)
    assert "ДОГОВОР ПОСТАВКИ № 291" in text
    assert "Сафин Рафис Рауфович" in text
    assert "027808041941" in text
    assert "2.1.5" not in text
    assert "{{" not in text


def test_company_intro_and_accountant_message():
    company = payload(
        buyer_type="ooo",
        buyer_name="ООО «МЕГАТРЕЙД»",
        representative_position="директора",
        representative_genitive="Шагаловой Анны Викторовны",
        representative_basis="Устава",
        buyer_inn="7017489758",
        buyer_kpp="701701001",
        buyer_ogrn="1207000009809",
    )
    assert "директора Шагаловой Анны Викторовны" in buyer_intro(company)
    assert "ООО «МЕГАТРЕЙД»" in accountant_message(company)
    assert "От Селеткова" in accountant_message(company)


def test_registry_allows_duplicate_number_and_exports(tmp_path):
    database = ContractsDB(tmp_path / "contracts.sqlite3")
    values = payload()
    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"
    database.add(values, 1, "Помощник", first)
    database.add(values, 2, "Другой помощник", second)
    assert len(database.by_number("291")) == 2
    assert database.search("027808041941")[0].creator_name == "Другой помощник"
    export = tmp_path / "registry.xlsx"
    assert database.export_excel(export) == 2
    workbook = load_workbook(export, read_only=True)
    assert workbook["Договоры"].max_row == 3
