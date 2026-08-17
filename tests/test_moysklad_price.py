import gzip
import json
from decimal import Decimal

from openpyxl import load_workbook

from moysklad_price import (
    BONUS_CATEGORIES,
    MoySkladClient,
    MoySkladError,
    build_price_from_moysklad,
    build_product_folder_mapping,
)
from prices_db import parse_price_file


class FakeMoySkladClient:
    def stores(self):
        return [
            {"name": name, "meta": {"href": f"https://api.moysklad.ru/api/remap/1.2/entity/store/{index}"}}
            for index, name in enumerate(("Мордор", "Годзибасы", "Жможики", "Другой"), 1)
        ]
    def current_availability(self, store_ids):
        assert store_ids == ("1", "2", "3")
        return [
            {
                "assortmentId": "available", "storeId": "1", "quantity": 7,
            },
            {
                "assortmentId": "available", "storeId": "3", "quantity": 2,
            },
            {
                "assortmentId": "reserved", "storeId": "1", "quantity": 0,
            },
            {
                "assortmentId": "other", "storeId": "4", "quantity": 20,
            },
        ]

    def assortment(self):
        prices = [
            {"priceType": {"name": "Дистр нал"}, "value": 10000},
            {"priceType": {"name": "Дистр безнал"}, "value": 11000},
            {"priceType": {"name": "от 50т.р. нал"}, "value": 12345},
            {"priceType": {"name": "от 50т.р. безнал"}, "value": 13579},
        ]
        return [
            {
                "meta": {
                    "href": "https://api.moysklad.ru/api/remap/1.2/entity/product/available",
                    "type": "product",
                },
                "name": "OGGO VLIQ Вишня",
                "pathName": "ЭС/Жидкости/OGGO VLIQ",
                "salePrices": prices,
            },
            {
                "meta": {
                    "href": "https://api.moysklad.ru/api/remap/1.2/entity/product/reserved",
                    "type": "product",
                },
                "name": "Полностью зарезервирован",
                "pathName": "ЭС/Жидкости/OGGO VLIQ",
                "salePrices": prices,
            },
        ]


class FakeGzipResponse:
    headers = {"Content-Encoding": "gzip"}

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    def read(self):
        return gzip.compress(json.dumps({"rows": []}).encode("utf-8"))


def test_api_requests_gzip_and_uses_get_only():
    captured = {}

    def opener(request, **kwargs):
        captured["method"] = request.get_method()
        captured["accept"] = request.get_header("Accept")
        captured["accept_encoding"] = request.get_header("Accept-encoding")
        return FakeGzipResponse()

    client = MoySkladClient("test-token", opener=opener)
    assert client.stores() == []
    assert captured == {
        "method": "GET",
        "accept": "application/json;charset=utf-8",
        "accept_encoding": "gzip",
    }


def test_current_availability_retries_without_store_filter_on_400():
    client = MoySkladClient("test-token", opener=lambda *args, **kwargs: None)
    calls = []

    def fake_get(url, params):
        calls.append(params)
        if len(calls) == 1:
            raise MoySkladError("bad filter", status=400)
        return [{"assortmentId": "item", "storeId": "1", "quantity": 3}]

    client._get_json = fake_get
    assert client.current_availability(("1", "2"))[0]["quantity"] == 3
    assert calls == [
        {"stockType": "quantity", "filter": "storeId=1,2"},
        {"stockType": "quantity"},
    ]


def test_builds_read_only_price_for_available_selected_stock(tmp_path):
    destination = tmp_path / "common.xlsx"
    result = build_price_from_moysklad(
        "test-token",
        destination,
        client=FakeMoySkladClient(),
    )

    assert result.item_count == 1
    assert result.group_count == 1
    parsed = parse_price_file(destination)
    assert parsed.item_count == 1
    item = parsed.groups[0].items[0]
    assert item.name == "OGGO VLIQ Вишня"
    assert item.cash == Decimal("123.45")
    assert item.cashless == Decimal("135.79")

    workbook = load_workbook(destination, data_only=True)
    sheet = workbook.active
    item_row = next(row for row in sheet.iter_rows(values_only=True) if row[0] == item.name)
    assert item_row[3:6] == (7, 0, 2)
    assert sheet.freeze_panes == "A4"
    workbook.close()


def test_fails_if_required_store_is_missing(tmp_path):
    client = FakeMoySkladClient()
    client.stores = lambda: []
    try:
        build_price_from_moysklad("test-token", tmp_path / "price.xlsx", client=client)
    except MoySkladError as error:
        assert "Не найдены склады" in str(error)
    else:
        raise AssertionError("Missing stores must stop price generation")


def test_builds_product_folder_mapping_from_moysklad(tmp_path):
    class FakeFolderClient:
        def product_folders(self):
            return [
                {"id": "root", "name": "ЭС", "pathName": ""},
                {"id": "liquids", "name": "Жидкости", "pathName": "ЭС"},
                {"id": "old", "name": "Архив", "pathName": "", "archived": True},
            ]

    destination = tmp_path / "folders.xlsx"
    count = build_product_folder_mapping(
        "test-token", destination, client=FakeFolderClient()
    )

    assert count == 1
    workbook = load_workbook(destination)
    sheet = workbook["Классификация папок"]
    values = list(sheet.iter_rows(min_row=2, values_only=True))
    assert values == [("ЭС/Жидкости", None, None, "liquids")]
    assert sheet.freeze_panes == "A2"
    assert len(sheet.data_validations.dataValidation) == 1
    guide_values = [cell.value for cell in workbook["Справочник"]["A"]][1:]
    assert guide_values == list(BONUS_CATEGORIES)
    workbook.close()
