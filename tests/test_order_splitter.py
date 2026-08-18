from decimal import Decimal

from openpyxl import Workbook, load_workbook

from order_splitter import OUTPUT_HEADERS, read_customer_order, split_customer_order


def customer_price(path):
    workbook = Workbook()
    sheet = workbook.active
    for _ in range(7):
        sheet.append([])
    sheet.append([
        "Код", "Наименование", "Ед.изм.", "от 50т.р. нал",
        "от 50т.р. безнал", "Итоговая цена", "Количество",
    ])
    sheet.append(["001", "Товар A", "шт", 100, 110, None, 10])
    sheet.append(["002", "Товар B", "шт", 200, 220, None, 5])
    sheet.append(["003", "Товар C", "шт", 300, 330, None, 2])
    sheet.append(["004", "Не заказан", "шт", 400, 440, None, None])
    workbook.save(path)
    workbook.close()


class FakeOrderClient:
    def stores(self):
        return [
            {
                "name": name,
                "meta": {"href": f"https://api.moysklad.ru/api/remap/1.2/entity/store/{number}"},
            }
            for number, name in enumerate(("Мордор", "Годзибасы", "Жможики"), 1)
        ]

    def current_availability(self, store_ids):
        assert store_ids == ("1", "2", "3")
        return [
            {"assortmentId": "a", "storeId": "1", "quantity": 4},
            {"assortmentId": "a", "storeId": "2", "quantity": 2},
            {"assortmentId": "a", "storeId": "3", "quantity": 6},
            {"assortmentId": "b", "storeId": "1", "quantity": 5},
        ]

    def assortment(self):
        return [
            {
                "code": "001",
                "name": "Товар A",
                "meta": {
                    "type": "product",
                    "href": "https://api.moysklad.ru/api/remap/1.2/entity/product/a",
                },
            },
            {
                "article": "002",
                "name": "Товар B",
                "meta": {
                    "type": "variant",
                    "href": "https://api.moysklad.ru/api/remap/1.2/entity/variant/b",
                },
            },
        ]


def test_reads_only_rows_with_customer_quantity(tmp_path):
    source = tmp_path / "customer.xlsx"
    customer_price(source)
    lines = read_customer_order(source)
    assert [line.code for line in lines] == ["001", "002", "003"]
    assert [line.quantity for line in lines] == [Decimal(10), Decimal(5), Decimal(2)]


def test_splits_order_with_priority_and_compact_outputs(tmp_path):
    source = tmp_path / "customer.xlsx"
    customer_price(source)
    result = split_customer_order(
        "token", source, tmp_path / "result", "Мордор", client=FakeOrderClient()
    )

    assert result.allocation_order == ("Мордор", "Жможики", "Годзибасы")
    assert [name for name, _ in result.files] == ["Мордор", "Жможики"]
    assert result.requested_quantity == Decimal(17)
    assert result.allocated_quantity == Decimal(15)
    assert result.shortage_quantity == Decimal(2)
    assert result.shortage_line_count == 1
    assert result.unknown_codes == ("003",)

    mordor = load_workbook(result.files[0][1], data_only=True)
    sheet = mordor.active
    assert tuple(cell.value for cell in sheet[1]) == OUTPUT_HEADERS
    assert sheet.max_column == 5
    assert sheet.max_row == 3
    assert sheet["A2"].value == "001"
    assert sheet["E2"].value == 4
    assert sheet["A3"].value == "002"
    assert sheet["E3"].value == 5
    mordor.close()

    zhmozhiki = load_workbook(result.files[1][1], data_only=True)
    assert zhmozhiki.active["A2"].value == "001"
    assert zhmozhiki.active["E2"].value == 6
    zhmozhiki.close()

    shortage = load_workbook(result.shortage_path, data_only=True)
    assert shortage.active.max_column == 5
    assert shortage.active["A2"].value == "003"
    assert shortage.active["E2"].value == 2
    shortage.close()
