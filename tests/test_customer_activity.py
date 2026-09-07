from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from customer_activity import (
    ActivityDatabase,
    build_activity_excel,
    build_monthly_activity_excel,
    client_parts,
    included_sales_channel,
    monthly_report_rows,
    report_rows,
    sync_shipments,
)


def demand(identifier, name, moment, channel="Валера", channel_id="v", amount=10000):
    return {
        "id": identifier, "name": identifier, "moment": moment,
        "updated": moment, "sum": amount,
        "agent": {"name": name, "meta": {"href": f"https://api.moysklad.ru/entity/counterparty/{identifier}"}},
        "salesChannel": {"name": channel, "meta": {"href": f"https://api.moysklad.ru/entity/saleschannel/{channel_id}"}},
    }


class FakeClient:
    def __init__(self, rows):
        self.rows = rows
        self.params = []

    def _rows(self, endpoint, params):
        self.params.append((endpoint, params))
        return self.rows


def test_client_names_are_grouped_by_prefix():
    assert client_parts("Is_getto/Шаров А.И.-5/Вологда")[0] == client_parts("Is getto/Литвиневский В.М.-5")[0]
    assert client_parts("(О) Vape Zone/НАЛ/СПб")[1] == "(О) Vape Zone"
    assert client_parts("Вячеслав   KAILAS {Ульяновск} | Нал")[0] == client_parts(
        "Вячеслав KAILAS {Ульяновск} | (ИП Ниськов А.Н.)"
    )[0]
    assert client_parts("Manager   Sistem Opt (Владимирская область)| (ИП Механиков Д. А.)")[0] == client_parts(
        "Manager Sistem Opt (Владимирская область) | (ИП Куликова А.С.)"
    )[0]
    assert client_parts("Joki   Липецк (ИП Нечаев В. А.)")[0] == client_parts(
        "Joki Липецк (ИП Лаврентьев   А.Н.)"
    )[0]
    assert client_parts("Joki Липецк (ИП Толчеев С.Г.)")[1] == "Joki Липецк"
    assert client_parts("Oleg Wow   Smoke {Самара} (ИП Исмаилов Р.Р.)")[0] == client_parts(
        "Oleg Wow Smoke {Самара} (ИП   Севостьянов Р.С.)"
    )[0]
    assert client_parts("Oleg Wow Smoke {Самара} (ИП Севостьянов Р.С.)")[1] == (
        "Oleg Wow Smoke {Самара}"
    )


def test_weekly_groups_and_latest_channel(tmp_path):
    db = ActivityDatabase(tmp_path / "activity.sqlite3")
    db.upsert_shipments([
        demand("old-return", "Return/ИП 1", "2026-06-01 10:00:00", "Андрей", "a"),
        demand("returned", "Return/ИП 2", "2026-09-03 10:00:00", "Валера", "v"),
        demand("new", "New shop/ИП", "2026-09-02 10:00:00", "Матвей", "m", 25000),
        demand("lost", "Lost/ИП", "2026-08-10 10:00:00", "Андрей", "a"),
    ])
    data = report_rows(db.shipments(), date(2026, 9, 7))
    assert [row["client"] for row in data["new"]] == ["New shop"]
    assert data["returned"][0]["client"] == "Return"
    assert data["returned"][0]["manager"] == "Валера"
    assert data["returned"][0]["gap"] == 94
    assert data["lost"][0]["category"] == "20–29 дней"
    assert len(data["previous"]["new"]) == 0
    assert len(data["previous"]["lost"]) == 2


def test_oggo_and_missing_latest_channels_are_excluded(tmp_path):
    assert included_sales_channel("Валера")
    assert not included_sales_channel("OGGO")
    assert not included_sales_channel("  oggo  ")
    assert not included_sales_channel("Без менеджера")
    assert not included_sales_channel("Розница")
    assert not included_sales_channel("  РОЗНИЦА ")
    assert not included_sales_channel("")

    db = ActivityDatabase(tmp_path / "excluded.sqlite3")
    db.upsert_shipments([
        demand("oggo", "OGGO client/ИП", "2026-08-01 10:00:00", "OGGO", "oggo"),
        demand("unassigned", "Unassigned/ИП", "2026-08-01 10:00:00", "Без менеджера", "none"),
        demand("blank", "No manager/ИП", "2026-08-01 10:00:00", "", "blank"),
        demand("regular", "Regular/ИП", "2026-08-01 10:00:00", "Валера", "v"),
    ])
    data = report_rows(db.shipments(), date(2026, 9, 7))
    assert [row["client"] for row in data["lost"]] == ["Regular"]
    assert set(data["dynamics"]["current"]) == {client_parts("Regular/ИП")[0]}


def test_monthly_report_tracks_customer_base_changes_and_excludes_retail(tmp_path):
    db = ActivityDatabase(tmp_path / "monthly.sqlite3")
    db.upsert_shipments([
        demand("return-old", "Returned/ИП 1", "2026-04-01 10:00:00", "Валера", "v"),
        demand("return-now", "Returned/ИП 2", "2026-08-10 10:00:00", "Валера", "v", 30000),
        demand("new", "New/ИП", "2026-08-05 10:00:00", "Валера", "v", 20000),
        demand("lost-one", "Lost/ИП 1", "2026-06-20 10:00:00", "Валера", "v"),
        demand("lost-two", "Lost/ИП 2", "2026-07-10 10:00:00", "Валера", "v"),
        demand("retail", "Retail/ИП", "2026-08-07 10:00:00", "Розница", "retail"),
    ])
    data = monthly_report_rows(db.shipments(), date(2026, 8, 1))
    assert [item["client"] for item in data["new"]] == ["New"]
    assert [item["client"] for item in data["returned"]] == ["Returned"]
    assert [item["client"] for item in data["became_inactive"]] == ["Lost"]
    assert {item["client"] for item in data["buyers"]} == {"New", "Returned"}

    destination = tmp_path / "monthly.xlsx"
    summary = build_monthly_activity_excel(destination, data)
    assert summary["active_start"] == 1
    assert summary["active_end"] == 2
    assert summary["active_change"] == 1
    assert summary["new"] == 1
    assert summary["returned"] == 1
    assert summary["became_inactive"] == 1
    workbook = load_workbook(destination)
    assert workbook.sheetnames == [
        "Итоги месяца", "Динамика", "Новые клиенты", "Вернувшиеся клиенты",
        "Стали неактивными", "Покупатели месяца",
    ]
    assert workbook["Итоги месяца"].freeze_panes == "A3"
    trend = workbook["Динамика"]
    assert trend.max_row == 6  # April through August.
    assert len(trend._charts) == 3
    assert [len(chart.ser) for chart in trend._charts] == [2, 3, 1]
    assert trend.freeze_panes == "A2"


def test_sync_uses_start_date_then_incremental_cursor(tmp_path):
    db = ActivityDatabase(tmp_path / "activity.sqlite3")
    client = FakeClient([demand("one", "Shop/ИП", "2026-08-01 10:00:00")])
    sync_shipments(db, client, datetime(2026, 9, 1, 3, 0))
    sync_shipments(db, client, datetime(2026, 9, 8, 3, 0))
    assert "moment>=2026-04-01 00:00:00" in client.params[0][1]["filter"]
    assert client.params[0][1]["limit"] == 100
    assert "updated>=" not in client.params[0][1]["filter"]
    assert "updated>=2026-09-01 02:59:55" in client.params[1][1]["filter"]


def test_inline_history_uses_local_grouped_shipments(tmp_path):
    db = ActivityDatabase(tmp_path / "inline-history.sqlite3")
    db.upsert_shipments([
        demand("one", "Vape Zone/ИП 1", "2026-06-01 10:00:00", amount=1000000),
        demand("two", "Vape Zone/ИП 2", "2026-07-01 10:00:00", amount=2000000),
        demand("three", "Vape Zone/ИП 2", "2026-08-01 10:00:00", amount=3000000),
    ])
    history = db.inline_client_history("Vape Zone/ИП 3", now=datetime(2026, 8, 11, 12, 0))
    assert history["days_since_last"] == 10
    assert history["recent_purchases"] == 3
    assert history["recent_revenue"] == 60000
    assert history["average_purchase"] == 20000
    assert history["expected_days"] == 30


def test_excel_has_dashboard_and_single_lost_sheet(tmp_path):
    destination = tmp_path / "report.xlsx"
    db = ActivityDatabase(tmp_path / "excel.sqlite3")
    db.upsert_shipments([demand("lost", "Lost/ИП", "2026-08-01 10:00:00")])
    data = report_rows(db.shipments(), date(2026, 9, 7))
    build_activity_excel(destination, data)
    workbook = load_workbook(destination)
    assert workbook.sheetnames == ["Дэшборд", "Помесячная аналитика", "Новые кенты", "Вернувшиеся кенты", "Кенты-потеряшки"]
    dashboard = workbook["Дэшборд"]
    assert len(dashboard._charts) == 2
    assert len(dashboard._charts[0].ser) == 2
    assert len(dashboard._charts[1].ser) == 1
    assert dashboard["A25"].value == "Как читать дэшборд"
    monthly_headers = [cell.value for cell in workbook["Помесячная аналитика"][1]]
    assert "Уникальные покупатели" in monthly_headers
    assert "Выручка на покупателя" in monthly_headers
    assert workbook["Помесячная аналитика"].max_row == 6  # April through August only.
    headers = [cell.value for cell in dashboard[20]]
    assert dashboard["F19"].value == "Клиенты, которые давно не заказывали"
    assert "Новые — предыдущая" not in headers
    assert "Вернувшиеся — предыдущая" not in headers
    assert "Активные неделю назад" not in headers
    assert dashboard.freeze_panes == "A3"
