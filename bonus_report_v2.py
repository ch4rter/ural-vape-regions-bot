from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


NAVY = "172331"
SLATE = "33475B"
BLUE = "4472C4"
TEAL = "18A999"
ORANGE = "F28E2B"
RED = "D9534F"
PALE = "EFF3F7"
WHITE = "FFFFFF"
TEXT = "263442"
MUTED = "6B7785"
MONEY = '#,##0.00 [$₽-ru-RU]'
DASHBOARD_MONEY = '#,##0 [$₽-ru-RU]'


def _amount(rows, attribute: str) -> Decimal:
    return sum((getattr(row, attribute) for row in rows), Decimal(0))


def _paint_range(sheet, cell_range: str, color: str) -> None:
    for row in sheet[cell_range]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=color)


def _card(sheet, start_col: int, label: str, value, number_format: str, color: str) -> None:
    end_col = start_col + 2
    _paint_range(sheet, f"{sheet.cell(4, start_col).coordinate}:{sheet.cell(6, end_col).coordinate}", WHITE)
    sheet.merge_cells(start_row=4, start_column=start_col, end_row=4, end_column=end_col)
    sheet.merge_cells(start_row=5, start_column=start_col, end_row=6, end_column=end_col)
    label_cell = sheet.cell(4, start_col, label.upper())
    label_cell.font = Font(name="Aptos", size=9, bold=True, color=MUTED)
    label_cell.alignment = Alignment(horizontal="center", vertical="center")
    value_cell = sheet.cell(5, start_col, value)
    value_cell.font = Font(name="Aptos Display", size=19, bold=True, color=color)
    value_cell.alignment = Alignment(horizontal="center", vertical="center")
    value_cell.number_format = number_format
    border = Side(style="thin", color="D8E0E8")
    for row in range(4, 7):
        for column in range(start_col, end_col + 1):
            cell = sheet.cell(row, column)
            cell.border = Border(
                left=border if column == start_col else None,
                right=border if column == end_col else None,
                top=border if row == 4 else None,
                bottom=border if row == 6 else None,
            )


def _add_table(sheet, reference: str, name: str) -> None:
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)


def build_bonus_workbook_v2(
    destination: Path,
    documents: list,
    categories: list[str],
    period_label: str,
    report_name: str,
) -> None:
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Дашборд"
    details = workbook.create_sheet("Отгрузки")
    data = workbook.create_sheet("Данные")

    total = _amount(documents, "total")
    paid = _amount(documents, "paid")
    debt = total - paid
    clients = {row.agent.casefold() for row in documents if row.agent and row.agent != "—"}
    average = total / len(documents) if documents else Decimal(0)
    category_totals = {
        category: sum((row.categories.get(category, Decimal(0)) for row in documents), Decimal(0))
        for category in categories
    }

    dashboard.sheet_view.showGridLines = False
    dashboard.freeze_panes = None
    dashboard.sheet_properties.pageSetUpPr.fitToPage = True
    dashboard.page_setup.fitToWidth = 1
    dashboard.page_setup.fitToHeight = 1
    dashboard.sheet_view.zoomScale = 72
    for column in range(1, 15):
        dashboard.column_dimensions[get_column_letter(column)].width = 12
    dashboard.row_dimensions[1].height = 34
    _paint_range(dashboard, "A1:N2", NAVY)
    dashboard.merge_cells("A1:N1")
    dashboard["A1"] = "ПРЕМИАЛЬНЫЙ ОТЧЁТ"
    dashboard["A1"].font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    dashboard["A1"].alignment = Alignment(horizontal="left", vertical="center")
    dashboard.merge_cells("A2:F2")
    dashboard["A2"] = f"Период: {period_label}"
    dashboard.merge_cells("G2:N2")
    dashboard["G2"] = f"Канал продаж: {report_name}"
    for coordinate in ("A2", "G2"):
        dashboard[coordinate].font = Font(name="Aptos", size=10, color="DCE5EE")
        dashboard[coordinate].alignment = Alignment(horizontal="left", vertical="center")

    cards = [
        (1, "Выручка", float(total), DASHBOARD_MONEY, BLUE),
        (4, "Оплачено", float(paid), DASHBOARD_MONEY, TEAL),
        (7, "Задолженность", float(debt), DASHBOARD_MONEY, RED if debt > 0 else TEAL),
        (10, "Оплата", float(paid / total if total else 0), "0.0%", ORANGE),
    ]
    for card in cards:
        _card(dashboard, *card)
    dashboard.merge_cells("A7:N7")
    dashboard["A7"] = (
        f"{len(documents)} отгрузок  ·  {len(clients)} контрагентов  ·  "
        f"средний чек {float(average):,.0f} ₽".replace(",", " ")
    )
    dashboard["A7"].font = Font(name="Aptos", size=10, color=MUTED)

    manager_rows: dict[str, list] = defaultdict(list)
    client_rows: dict[tuple[str, str], list] = defaultdict(list)
    for document in documents:
        manager = document.sales_channel or "Без канала"
        manager_rows[manager].append(document)
        client_rows[(manager, document.agent)].append(document)

    data.append(["Категория", "Сумма"])
    for category, amount in sorted(category_totals.items(), key=lambda item: (-item[1], item[0].casefold())):
        data.append([category, float(amount)])
    category_end = data.max_row
    manager_start = data.max_row + 3
    data.cell(manager_start, 1, "Менеджер")
    data.cell(manager_start, 2, "Выручка")
    data.cell(manager_start, 3, "Оплачено")
    data.cell(manager_start, 4, "Задолженность")
    data.cell(manager_start, 5, "Контрагенты")
    data.cell(manager_start, 6, "Отгрузки")
    manager_order = sorted(manager_rows, key=lambda key: (-_amount(manager_rows[key], "total"), key.casefold()))
    for manager in manager_order:
        rows = manager_rows[manager]
        manager_total = _amount(rows, "total")
        manager_paid = _amount(rows, "paid")
        data.append([
            manager, float(manager_total), float(manager_paid), float(manager_total - manager_paid),
            len({row.agent.casefold() for row in rows if row.agent}), len(rows),
        ])
    manager_end = data.max_row
    client_start = data.max_row + 3
    data.cell(client_start, 1, "Контрагент")
    data.cell(client_start, 2, "Менеджер")
    data.cell(client_start, 3, "Сумма")
    data.cell(client_start, 4, "Оплачено")
    data.cell(client_start, 5, "Задолженность")
    ordered_clients = sorted(client_rows.values(), key=lambda rows: -_amount(rows, "total"))
    for rows in ordered_clients:
        client_total = _amount(rows, "total")
        client_paid = _amount(rows, "paid")
        data.append([
            rows[0].agent, rows[0].sales_channel or "Без канала", float(client_total),
            float(client_paid), float(client_total - client_paid),
        ])
    client_end = data.max_row
    payment_start = data.max_row + 3
    data.cell(payment_start, 1, "Статус")
    data.cell(payment_start, 2, "Сумма")
    data.append(["Оплачено", float(paid)])
    data.append(["Задолженность", float(max(debt, Decimal(0)))])

    premium_categories = list(categories)
    manager_headers = [
        "Менеджер", "Клиенты", "Отгрузки", "Выручка", "Оплачено", "Долг", "% оплаты",
        *premium_categories,
    ]
    manager_last_column = len(manager_headers)
    manager_last_letter = get_column_letter(manager_last_column)
    _paint_range(dashboard, f"A8:{manager_last_letter}8", SLATE)
    dashboard.merge_cells(f"A8:{manager_last_letter}8")
    dashboard["A8"] = "РЕЗУЛЬТАТЫ МЕНЕДЖЕРОВ И ПРЕМИАЛЬНЫХ НАПРАВЛЕНИЙ"
    dashboard["A8"].font = Font(name="Aptos", size=11, bold=True, color=WHITE)
    for column, header in enumerate(manager_headers, 1):
        cell = dashboard.cell(9, column, header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    manager_first_row = 10
    for row_number, manager in enumerate(manager_order, manager_first_row):
        rows = manager_rows[manager]
        manager_total = _amount(rows, "total")
        manager_paid = _amount(rows, "paid")
        manager_clients = {row.agent.casefold() for row in rows if row.agent and row.agent != "—"}
        values = [
            manager, len(manager_clients), len(rows), float(manager_total), float(manager_paid),
            float(manager_total - manager_paid), float(manager_paid / manager_total if manager_total else 0),
            *(float(sum((item.categories.get(category, Decimal(0)) for item in rows), Decimal(0)))
              for category in premium_categories),
        ]
        for column, value in enumerate(values, 1):
            dashboard.cell(row_number, column, value)

    manager_total_row = manager_first_row + len(manager_order)
    totals = [
        "ИТОГО", len(clients), len(documents), float(total), float(paid), float(debt),
        float(paid / total if total else 0),
        *(float(category_totals[category]) for category in premium_categories),
    ]
    for column, value in enumerate(totals, 1):
        dashboard.cell(manager_total_row, column, value)
    for row_number in range(manager_first_row, manager_total_row + 1):
        if row_number < manager_total_row and row_number % 2 == 0:
            for column in range(1, manager_last_column + 1):
                dashboard.cell(row_number, column).fill = PatternFill("solid", fgColor=PALE)
        for column in range(4, manager_last_column + 1):
            dashboard.cell(row_number, column).number_format = DASHBOARD_MONEY
        dashboard.cell(row_number, 7).number_format = "0.0%"
    for column in range(1, manager_last_column + 1):
        dashboard.cell(manager_total_row, column).fill = PatternFill("solid", fgColor="DDEBF7")
        dashboard.cell(manager_total_row, column).font = Font(name="Aptos", bold=True, color=TEXT)
    dashboard.column_dimensions["A"].width = 18
    dashboard.column_dimensions["B"].width = 8
    dashboard.column_dimensions["C"].width = 9
    for column in range(4, manager_last_column + 1):
        dashboard.column_dimensions[get_column_letter(column)].width = 16
    dashboard.row_dimensions[9].height = 34

    lower_row = manager_total_row + 2
    _paint_range(dashboard, f"A{lower_row}:F{lower_row}", SLATE)
    dashboard.merge_cells(start_row=lower_row, start_column=1, end_row=lower_row, end_column=6)
    dashboard.cell(lower_row, 1, "ТОП-5 КОНТРАГЕНТОВ")
    dashboard.cell(lower_row, 1).font = Font(name="Aptos", bold=True, color=WHITE)
    client_headers = ["Контрагент", "Менеджер", "Отгрузки", "Выручка", "Оплачено", "Долг"]
    for column, header in enumerate(client_headers, 1):
        cell = dashboard.cell(lower_row + 1, column, header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
    for row_number, rows in enumerate(ordered_clients[:5], lower_row + 2):
        client_total = _amount(rows, "total")
        client_paid = _amount(rows, "paid")
        values = [
            rows[0].agent, rows[0].sales_channel or "Без канала", len(rows),
            float(client_total), float(client_paid), float(client_total - client_paid),
        ]
        for column, value in enumerate(values, 1):
            dashboard.cell(row_number, column, value)
        for column in (4, 5, 6):
            dashboard.cell(row_number, column).number_format = DASHBOARD_MONEY

    detail_headers = [
        "Тип", "Дата", "Документ", "Контрагент", "Статус", "Канал продаж",
        "Сумма документа", "Оплачено", "Задолженность", *categories,
        "Контрольное расхождение",
    ]
    details.append(["ДЕТАЛИЗАЦИЯ ОТГРУЗОК", period_label, report_name])
    details.append([])
    details.append(detail_headers)
    for document in documents:
        categorized = sum(document.categories.values(), Decimal(0))
        details.append([
            document.kind, document.moment[:10], document.name, document.agent,
            document.state, document.sales_channel, float(document.total), float(document.paid),
            float(document.total - document.paid),
            *(float(document.categories.get(category, Decimal(0))) for category in categories),
            float(document.total - categorized),
        ])
    details.sheet_view.showGridLines = False
    details.freeze_panes = "A4"
    details.row_dimensions[1].height = 28
    for cell in details[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos Display", bold=True, color=WHITE, size=13)
    if details.max_row >= 4:
        _add_table(details, f"A3:{details.cell(details.max_row, len(detail_headers)).coordinate}", "BonusShipmentsV2")
    widths = [12, 13, 18, 42, 22, 22, 18, 18, 18] + [22] * len(categories) + [22]
    for index, width in enumerate(widths, 1):
        details.column_dimensions[details.cell(1, index).column_letter].width = width
    for row in range(4, details.max_row + 1):
        for column in range(7, len(detail_headers) + 1):
            details.cell(row, column).number_format = MONEY

    data.sheet_view.showGridLines = False
    data.column_dimensions["A"].width = 38
    data.column_dimensions["B"].width = 22
    for column in ("C", "D", "E", "F"):
        data.column_dimensions[column].width = 19
    if category_end >= 2:
        _add_table(data, f"A1:B{category_end}", "BonusCategoriesV2")
    if manager_end > manager_start:
        _add_table(data, f"A{manager_start}:F{manager_end}", "BonusManagersV2")
    if client_end > client_start:
        _add_table(data, f"A{client_start}:E{client_end}", "BonusClientsV2")
    _add_table(data, f"A{payment_start}:B{payment_start + 2}", "BonusPaymentsV2")
    for row in range(2, data.max_row + 1):
        for column in range(2, min(data.max_column, 5) + 1):
            if isinstance(data.cell(row, column).value, (int, float)):
                data.cell(row, column).number_format = MONEY
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()
