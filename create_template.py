from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo


path = Path(__file__).resolve().parent / "managers.xlsx"
if path.exists():
    raise SystemExit(f"Файл уже существует и не был перезаписан: {path}")

workbook = Workbook()
sheet = workbook.active
sheet.title = "Менеджеры"
sheet.append(["Название", "Менеджер"])
sheet.append(["Тамбов", "Андрей"])
sheet.append(["Тамбовская область", "Андрей"])
sheet.append(["Минск", "Елена"])
sheet.append(["Алматы", "Сергей"])
sheet.freeze_panes = "A2"
sheet.column_dimensions["A"].width = 34
sheet.column_dimensions["B"].width = 24
for cell in sheet[1]:
    cell.font = Font(bold=True)

table = Table(displayName="Managers", ref=f"A1:B{sheet.max_row}")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
sheet.add_table(table)
workbook.save(path)
print(f"Создан шаблон: {path}")
