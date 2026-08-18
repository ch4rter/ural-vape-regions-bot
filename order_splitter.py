"""Read-only distribution of a filled customer price list across warehouses."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from moysklad_price import (
    MoySkladClient,
    MoySkladError,
    _available_stock,
    _canonical_href,
    _select_stores,
)


ORDER_STORES = ("Мордор", "Годзибасы", "Жможики")
OUTPUT_HEADERS = (
    "Код товара",
    "Наименование",
    "Цена нал",
    "Цена безнал",
    "Количество",
)


@dataclass(frozen=True)
class CustomerOrderLine:
    code: str
    name: str
    cash_price: Decimal
    cashless_price: Decimal
    quantity: Decimal


@dataclass(frozen=True)
class AllocatedOrderLine:
    source: CustomerOrderLine
    quantity: Decimal


@dataclass(frozen=True)
class OrderSplitResult:
    files: tuple[tuple[str, Path], ...]
    shortage_path: Path | None
    allocation_order: tuple[str, ...]
    line_count: int
    requested_quantity: Decimal
    allocated_quantity: Decimal
    shortage_quantity: Decimal
    shortage_line_count: int
    unknown_codes: tuple[str, ...]


def _normalized_header(value: object) -> str:
    return re.sub(r"[^a-zа-яё0-9]+", "", str(value or "").casefold())


def _decimal(value: object, *, field: str, row_number: int) -> Decimal:
    if value is None or isinstance(value, bool):
        raise MoySkladError(f"Строка {row_number}: не заполнено поле «{field}».")
    text = str(value).strip().replace(" ", "").replace(",", ".")
    try:
        result = Decimal(text)
    except (InvalidOperation, ValueError) as error:
        raise MoySkladError(
            f"Строка {row_number}: в поле «{field}» указано некорректное число «{value}»."
        ) from error
    if not result.is_finite():
        raise MoySkladError(f"Строка {row_number}: некорректное значение поля «{field}».")
    return result


def _code(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_customer_order(path: Path) -> tuple[CustomerOrderLine, ...]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise MoySkladError("Не удалось открыть Excel-файл заказа.") from error
    try:
        sheet = workbook.worksheets[0]
        aliases = {
            "code": {"код", "кодтовара"},
            "name": {"наименование", "названиетовара"},
            "cash": {"от50трнал", "ценанал", "нал"},
            "cashless": {"от50трбезнал", "ценабезнал", "безнал"},
            "quantity": {"количество", "колво"},
        }
        header_row = 0
        columns: dict[str, int] = {}
        for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(30, sheet.max_row)), 1):
            found: dict[str, int] = {}
            for column, cell in enumerate(row, 1):
                normalized = _normalized_header(cell.value)
                for key, values in aliases.items():
                    if normalized in values:
                        found[key] = column
            if set(found) == set(aliases):
                header_row = row_number
                columns = found
                break
        if not header_row:
            raise MoySkladError(
                "В первых 30 строках не найдены колонки «Код», «Наименование», "
                "«от 50т.р. нал», «от 50т.р. безнал» и «Количество»."
            )

        by_code: dict[str, CustomerOrderLine] = {}
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1
        ):
            quantity_value = values[columns["quantity"] - 1]
            if quantity_value in (None, ""):
                continue
            quantity = _decimal(quantity_value, field="Количество", row_number=row_number)
            if quantity == 0:
                continue
            if quantity < 0:
                raise MoySkladError(f"Строка {row_number}: количество не может быть отрицательным.")
            code = _code(values[columns["code"] - 1])
            name = str(values[columns["name"] - 1] or "").strip()
            if not code or not name:
                raise MoySkladError(
                    f"Строка {row_number}: у заказанной позиции должны быть код и наименование."
                )
            cash = _decimal(values[columns["cash"] - 1], field="Цена нал", row_number=row_number)
            cashless = _decimal(
                values[columns["cashless"] - 1], field="Цена безнал", row_number=row_number
            )
            if cash < 0 or cashless < 0:
                raise MoySkladError(f"Строка {row_number}: цена не может быть отрицательной.")
            key = code.casefold()
            current = by_code.get(key)
            if current:
                if (current.name, current.cash_price, current.cashless_price) != (name, cash, cashless):
                    raise MoySkladError(
                        f"Код «{code}» встречается несколько раз с разными данными."
                    )
                by_code[key] = CustomerOrderLine(
                    code, name, cash, cashless, current.quantity + quantity
                )
            else:
                by_code[key] = CustomerOrderLine(code, name, cash, cashless, quantity)
        if not by_code:
            raise MoySkladError("В колонке «Количество» нет заказанных позиций.")
        return tuple(by_code.values())
    finally:
        workbook.close()


def _assortment_by_requested_code(
    assortment: list[dict], requested: set[str]
) -> dict[str, str]:
    result: dict[str, str] = {}
    duplicates: set[str] = set()
    for item in assortment:
        if item.get("archived") is True:
            continue
        if str(item.get("meta", {}).get("type", "")) not in {"product", "variant"}:
            continue
        assortment_id = _canonical_href(
            str(item.get("meta", {}).get("href", ""))
        ).rsplit("/", 1)[-1]
        if not assortment_id:
            continue
        identifiers = {
            _code(item.get("code")).casefold(),
            _code(item.get("article")).casefold(),
        } - {""}
        for identifier in identifiers & requested:
            if identifier in result and result[identifier] != assortment_id:
                duplicates.add(identifier)
            else:
                result[identifier] = assortment_id
    if duplicates:
        display = ", ".join(sorted(duplicates)[:10])
        raise MoySkladError(
            "В МоемСкладе нескольким товарам назначен один код: " + display + "."
        )
    return result


def _write_order_file(
    path: Path,
    lines: list[AllocatedOrderLine],
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Заказ"
    sheet.append(OUTPUT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for item in lines:
        sheet.append([
            item.source.code,
            item.source.name,
            float(item.source.cash_price),
            float(item.source.cashless_price),
            float(item.quantity),
        ])
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 3).number_format = '#,##0.00'
        sheet.cell(row, 4).number_format = '#,##0.00'
        sheet.cell(row, 5).number_format = '0.###'
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:E{sheet.max_row}"
    for column, width in zip("ABCDE", (16, 75, 16, 18, 14)):
        sheet.column_dimensions[column].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def split_customer_order(
    token: str,
    source: Path,
    destination_dir: Path,
    priority_store: str,
    *,
    client: MoySkladClient | None = None,
) -> OrderSplitResult:
    if priority_store not in ORDER_STORES:
        raise MoySkladError("Выбран неизвестный приоритетный склад.")
    lines = read_customer_order(source)
    client = client or MoySkladClient(token)
    selected_stores = _select_stores(client.stores(), ORDER_STORES)
    store_ids = tuple(
        _canonical_href(str(store.get("meta", {}).get("href", ""))).rsplit("/", 1)[-1]
        for store in selected_stores
    )
    if any(not store_id for store_id in store_ids):
        raise MoySkladError("Для разрешённых складов не получены API-идентификаторы.")

    stock = _available_stock(client.current_availability(store_ids), store_ids)
    requested_codes = {line.code.casefold() for line in lines}
    assortment_ids = _assortment_by_requested_code(client.assortment(), requested_codes)
    store_index = {name: index for index, name in enumerate(ORDER_STORES)}

    primary_index = store_index[priority_store]
    remaining_after_primary: dict[str, Decimal] = {}
    for line in lines:
        assortment_id = assortment_ids.get(line.code.casefold(), "")
        balances = stock.get(assortment_id, (Decimal(0),) * len(ORDER_STORES))
        remaining_after_primary[line.code.casefold()] = max(
            Decimal(0), line.quantity - balances[primary_index]
        )

    secondary = [name for name in ORDER_STORES if name != priority_store]
    secondary.sort(
        key=lambda name: (
            -sum(
                min(
                    remaining_after_primary[line.code.casefold()],
                    stock.get(
                        assortment_ids.get(line.code.casefold(), ""),
                        (Decimal(0),) * len(ORDER_STORES),
                    )[store_index[name]],
                )
                for line in lines
            ),
            ORDER_STORES.index(name),
        )
    )
    allocation_order = (priority_store, *secondary)
    allocations: dict[str, list[AllocatedOrderLine]] = {name: [] for name in ORDER_STORES}
    shortages: list[AllocatedOrderLine] = []
    unknown_codes = []
    for line in lines:
        code_key = line.code.casefold()
        assortment_id = assortment_ids.get(code_key, "")
        if not assortment_id:
            unknown_codes.append(line.code)
        balances = stock.get(assortment_id, (Decimal(0),) * len(ORDER_STORES))
        remaining = line.quantity
        for store_name in allocation_order:
            available = balances[store_index[store_name]]
            allocated = min(remaining, available)
            if allocated > 0:
                allocations[store_name].append(AllocatedOrderLine(line, allocated))
                remaining -= allocated
            if remaining <= 0:
                break
        if remaining > 0:
            shortages.append(AllocatedOrderLine(line, remaining))

    destination_dir.mkdir(parents=True, exist_ok=True)
    stamp = date.today().strftime("%d.%m.%Y")
    files = []
    for store_name in allocation_order:
        store_lines = allocations[store_name]
        if not store_lines:
            continue
        path = destination_dir / f"Заказ {store_name} {stamp}.xlsx"
        _write_order_file(path, store_lines)
        files.append((store_name, path))
    shortage_path = None
    if shortages:
        shortage_path = destination_dir / f"Не распределено {stamp}.xlsx"
        _write_order_file(shortage_path, shortages)

    requested_quantity = sum((line.quantity for line in lines), Decimal(0))
    shortage_quantity = sum((line.quantity for line in shortages), Decimal(0))
    return OrderSplitResult(
        files=tuple(files),
        shortage_path=shortage_path,
        allocation_order=allocation_order,
        line_count=len(lines),
        requested_quantity=requested_quantity,
        allocated_quantity=requested_quantity - shortage_quantity,
        shortage_quantity=shortage_quantity,
        shortage_line_count=len(shortages),
        unknown_codes=tuple(sorted(unknown_codes, key=str.casefold)),
    )
