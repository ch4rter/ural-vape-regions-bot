"""Editable commercial grouping rules for inline inventory cards."""
from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from prices_db import ItemSummary, normalize_price_text


@dataclass(frozen=True)
class InventoryGroupingRule:
    source_group: str
    source_category: str
    card_name: str
    line_name: str
    kind: str
    unit: str


def rule_key(group: str, category: str) -> str:
    return normalize_price_text(group) + "|" + normalize_price_text(category)


def load_grouping(path: Path) -> dict[str, InventoryGroupingRule]:
    if not path.is_file():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    result = {}
    for value in data.get("rules", []):
        rule = InventoryGroupingRule(**value)
        result[rule_key(rule.source_group, rule.source_category)] = rule
    return result


def export_grouping(path: Path, catalog: list[ItemSummary], current: dict[str, InventoryGroupingRule]) -> int:
    groups = sorted({(x.group_name, x.category_name) for x in catalog}, key=lambda x: normalize_price_text(x[0]))
    wb = Workbook(); ws = wb.active; ws.title = "Правила"
    ws.append(["Текущая группа", "Категория", "Общая карточка", "Название линейки", "Тип ассортимента", "Единица"])
    for group, category in groups:
        old = current.get(rule_key(group, category))
        ws.append([group, category, old.card_name if old else "", old.line_name if old else group,
                   old.kind if old else category, old.unit if old else "шт."])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="374151")
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDEF", (38, 28, 32, 42, 24, 16)): ws.column_dimensions[col].width = width
    guide = wb.create_sheet("Инструкция")
    guide.append(["Заполняйте «Общая карточка» только у групп, которые нужно объединить."])
    guide.append(["Одинаковая общая карточка собирает несколько линеек в один inline-результат."])
    guide.append(["Тип: жидкости/конструкторы/одноразки → вкусы; устройства → цвета; остальное → варианты."])
    path.parent.mkdir(parents=True, exist_ok=True); wb.save(path); return len(groups)


def save_grouping(source: Path, destination: Path) -> int:
    wb = load_workbook(source, read_only=True, data_only=True); ws = wb.worksheets[0]
    headers = {str(c.value or "").strip().casefold(): i for i, c in enumerate(ws[1])}
    required = ["текущая группа", "категория", "общая карточка", "название линейки", "тип ассортимента", "единица"]
    if any(x not in headers for x in required):
        raise ValueError("В таблице изменены или удалены обязательные колонки.")
    rules = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        get = lambda name: str(row[headers[name]] or "").strip()
        card = get("общая карточка")
        if not card: continue
        rule = InventoryGroupingRule(get("текущая группа"), get("категория"), card,
            get("название линейки") or get("текущая группа"), get("тип ассортимента") or get("категория"),
            get("единица") or "шт.")
        if not rule.source_group: raise ValueError("Обнаружена строка без текущей группы.")
        rules.append(rule)
    wb.close()
    payload = {"version": 1, "rules": [r.__dict__ for r in rules]}
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(".tmp"); temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"); temporary.replace(destination)
    return len(rules)
