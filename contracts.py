"""Contract registry and DOCX generation for supported suppliers."""

from __future__ import annotations

import json
import re
import sqlite3
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


SUPPLIERS = {
    "seletkov": {
        "short": "ИП Селетков В.Г.",
        "accountant": "Селеткова",
        "template": "contract_seletkov.docx",
    },
    "shmidt": {
        "short": "ИП Шмидт А.В.",
        "accountant": "Шмидта",
        "template": "contract_shmidt.docx",
    },
}
DEFAULT_SUPPLIER = "seletkov"


def supplier(payload: dict) -> dict:
    return SUPPLIERS.get(payload.get("supplier_key", DEFAULT_SUPPLIER), SUPPLIERS[DEFAULT_SUPPLIER])


@dataclass(frozen=True)
class ContractRecord:
    id: int
    number: str
    contract_date: str
    supplier_key: str
    buyer_type: str
    buyer_name: str
    buyer_inn: str
    edo_id: str
    creator_id: int
    creator_name: str
    file_path: str
    payload: dict
    created_at: str


class ContractsDB:
    def __init__(self, path: Path):
        self.path = path
        path.parent.mkdir(parents=True, exist_ok=True)
        with self._connect() as db:
            db.executescript(
                """
                CREATE TABLE IF NOT EXISTS contracts (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    number TEXT NOT NULL,
                    contract_date TEXT NOT NULL,
                    buyer_type TEXT NOT NULL,
                    buyer_name TEXT NOT NULL,
                    buyer_inn TEXT NOT NULL,
                    edo_id TEXT NOT NULL DEFAULT '',
                    creator_id INTEGER NOT NULL,
                    creator_name TEXT NOT NULL DEFAULT '',
                    file_path TEXT NOT NULL,
                    payload_json TEXT NOT NULL,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                );
                CREATE INDEX IF NOT EXISTS contracts_number_idx ON contracts(number);
                CREATE INDEX IF NOT EXISTS contracts_inn_idx ON contracts(buyer_inn);
                """
            )
            columns = {row[1] for row in db.execute("PRAGMA table_info(contracts)")}
            if "supplier_key" not in columns:
                db.execute(
                    "ALTER TABLE contracts ADD COLUMN supplier_key TEXT NOT NULL DEFAULT 'seletkov'"
                )

    def _connect(self):
        db = sqlite3.connect(self.path, timeout=30)
        db.row_factory = sqlite3.Row
        return db

    @staticmethod
    def _record(row: sqlite3.Row) -> ContractRecord:
        return ContractRecord(
            row["id"], row["number"], row["contract_date"], row["supplier_key"], row["buyer_type"],
            row["buyer_name"], row["buyer_inn"], row["edo_id"], row["creator_id"],
            row["creator_name"], row["file_path"], json.loads(row["payload_json"]),
            row["created_at"],
        )

    def by_number(self, number: str, supplier_key: str | None = None) -> list[ContractRecord]:
        with self._connect() as db:
            if supplier_key:
                rows = db.execute(
                    "SELECT * FROM contracts WHERE number=? AND supplier_key=? ORDER BY id DESC",
                    (number.strip(), supplier_key),
                ).fetchall()
            else:
                rows = db.execute(
                    "SELECT * FROM contracts WHERE number=? ORDER BY id DESC", (number.strip(),)
                ).fetchall()
        return [self._record(row) for row in rows]

    def recent(self, limit: int = 20) -> list[ContractRecord]:
        with self._connect() as db:
            rows = db.execute(
                "SELECT * FROM contracts ORDER BY id DESC LIMIT ?", (limit,)
            ).fetchall()
        return [self._record(row) for row in rows]

    def search(self, query: str, limit: int = 20) -> list[ContractRecord]:
        value = query.strip()
        like = f"%{value}%"
        with self._connect() as db:
            rows = db.execute(
                """SELECT * FROM contracts
                   WHERE number LIKE ? OR buyer_inn LIKE ? OR buyer_name LIKE ?
                   ORDER BY id DESC LIMIT ?""",
                (like, like, like, limit),
            ).fetchall()
        return [self._record(row) for row in rows]

    def add(self, payload: dict, creator_id: int, creator_name: str, file_path: Path) -> int:
        with self._connect() as db:
            cursor = db.execute(
                """INSERT INTO contracts(
                       number,contract_date,supplier_key,buyer_type,buyer_name,buyer_inn,edo_id,
                       creator_id,creator_name,file_path,payload_json
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    payload["contract_number"], payload["contract_date"],
                    payload.get("supplier_key", DEFAULT_SUPPLIER), payload["buyer_type"],
                    payload["buyer_name"], payload["buyer_inn"],
                    payload.get("edo_id", ""), creator_id, creator_name, str(file_path),
                    json.dumps(payload, ensure_ascii=False),
                ),
            )
            return int(cursor.lastrowid)

    def export_excel(self, destination: Path) -> int:
        rows = self.recent(100000)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Договоры"
        sheet.append([
            "Номер", "Дата договора", "Поставщик", "Тип", "Покупатель", "ИНН", "ЭДО",
            "Создал", "Telegram ID", "Дата формирования", "Файл",
        ])
        for record in rows:
            sheet.append([
                record.number, record.contract_date, SUPPLIERS.get(record.supplier_key, SUPPLIERS[DEFAULT_SUPPLIER])["short"], record.buyer_type.upper(),
                record.buyer_name, record.buyer_inn, record.edo_id,
                record.creator_name, record.creator_id, record.created_at, record.file_path,
            ])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="374151")
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:K{max(1, sheet.max_row)}"
        for column, width in {"A": 13, "B": 15, "C": 23, "D": 10, "E": 42, "F": 18, "G": 42, "H": 24, "I": 16, "J": 21, "K": 55}.items():
            sheet.column_dimensions[column].width = width
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
        return len(rows)

    def backup_to(self, destination: Path) -> None:
        destination.parent.mkdir(parents=True, exist_ok=True)
        with self._connect() as source, sqlite3.connect(destination) as target:
            source.backup(target)


def digits(value: str) -> str:
    return re.sub(r"\D", "", value)


def safe_file_part(value: str, limit: int = 70) -> str:
    result = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", " ", value)
    return " ".join(result.split())[:limit].strip(" .") or "Контрагент"


def buyer_intro(payload: dict) -> str:
    name = payload["buyer_name"].strip()
    representative = payload["representative_genitive"].strip()
    if payload["buyer_type"] == "ip":
        return (
            f"индивидуальный предприниматель {name}, именуемый в дальнейшем ПОКУПАТЕЛЬ, "
            f"в лице {representative}, действующего на основании свидетельства о регистрации "
            f"№ {payload['buyer_ogrn']}"
        )
    position = payload["representative_position"].strip()
    return (
        f"{name}, именуемое в дальнейшем ПОКУПАТЕЛЬ, в лице {position} {representative}, "
        f"действующего на основании {payload['representative_basis']}, "
        f"ОГРН {payload['buyer_ogrn']}"
    )


def buyer_requisites(payload: dict) -> str:
    buyer_name = (
        f"ИП {payload['buyer_name']}" if payload["buyer_type"] == "ip"
        else payload["buyer_name"]
    )
    values = [
        "ПОКУПАТЕЛЬ:", buyer_name,
        f"ИНН: {payload['buyer_inn']}",
    ]
    if payload.get("buyer_kpp"):
        values.append(f"КПП: {payload['buyer_kpp']}")
    values.extend([
        f"ОГРН{'ИП' if payload['buyer_type'] == 'ip' else ''}: {payload['buyer_ogrn']}",
        f"Адрес: {payload['buyer_address']}",
        f"Р/счет: {payload['bank_account']}",
        f"К/счет: {payload['correspondent_account']}",
        f"БИК: {payload['bik']}",
        f"Банк: {payload['bank_name']}",
    ])
    return "\n".join(values)


def buyer_short_name(payload: dict) -> str:
    if payload["buyer_type"] != "ip":
        return payload["buyer_name"].strip()
    parts = payload["buyer_name"].split()
    if not parts:
        return "ИП"
    initials = "".join(f"{part[0].upper()}." for part in parts[1:3] if part)
    return f"ИП {parts[0]} {initials}".strip()


def accountant_message(payload: dict) -> str:
    return (
        f"№{payload['contract_number']} от {payload['contract_date']}\n\n"
        f"От {supplier(payload)['accountant']}\n\n"
        f"{buyer_short_name(payload)}\n\n"
        f"{payload['buyer_inn']}\n\n"
        f"{payload['edo_id']}"
    )


def _replace_paragraph(paragraph, replacements: dict[str, str]) -> None:
    original = "".join(run.text for run in paragraph.runs)
    if not original or not any(key in original for key in replacements):
        return
    updated = original
    for key, value in replacements.items():
        updated = updated.replace(key, value)
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = updated


def render_contract(template: Path, destination: Path, payload: dict) -> None:
    document = Document(template)
    replacements = {
        "{{CONTRACT_NUMBER}}": payload["contract_number"],
        "{{CONTRACT_DATE}}": payload["contract_date"],
        "{{BUYER_INTRO}}": buyer_intro(payload),
        "{{BUYER_REQUISITES}}": buyer_requisites(payload),
    }
    for paragraph in document.paragraphs:
        _replace_paragraph(paragraph, replacements)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_paragraph(paragraph, replacements)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def create_contract_file(
    template: Path, directory: Path, payload: dict
) -> Path:
    date_part = datetime.strptime(payload["contract_date"], "%d.%m.%Y").strftime("%Y-%m-%d")
    filename = (
        f"Договор №{safe_file_part(payload['contract_number'], 20)} — "
        f"{safe_file_part(payload['buyer_name'])} — {date_part}.docx"
    )
    destination = directory / filename
    suffix = 2
    while destination.exists():
        destination = directory / f"{Path(filename).stem} ({suffix}).docx"
        suffix += 1
    render_contract(template, destination, payload)
    return destination
